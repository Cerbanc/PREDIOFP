"""Agenda de turnos del día en Excel, con el mismo diseño que la planilla de papel que se usa en el predio.

    C:\\Predio\\dias\\2026-10-15\\2026-10-15_agenda.xlsx

Una hoja ("Agenda"), apaisada y ajustada para imprimir en una sola hoja A4: arriba el día, después el bloque de FÚTBOL
(una fila por hora, cancha 1 a la izquierda, cancha 2 a la derecha y en el medio las dos casillas «5_7» que dicen qué cancha se jugó),
el bloque de PÁDEL (cancha 1, la columna del medio con los horarios y cancha 2) y abajo OBSERVACIONES con las notas de los turnos.
Las columnas «$» y «MP» quedan vacías: se completan a mano (si alguien las completa en el Excel, se conservan cuando la agenda se vuelve a armar).

Se hace para hoy y los 7 días siguientes. Se vuelve a armar sola cuando cambian turnos, turnos fijos o canchas (ArchivosDiarios la avisa en
cada guardado) y al empezar cada día; si el archivo no cambió respecto de lo último escrito, no se toca. Si está abierto en Excel se reintenta solo.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import logging
import re
import sqlite3
import threading
import time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from .diario import FORMATO_PLATA, ArchivoEnUso, _escribir_atomico, _limpiar_hoja, _san
from .util import hm

log = logging.getLogger("predio.agenda")

DIAS_ADELANTE = 7                                   # además de hoy
DIAS_MAYUS = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]
CLAVES_FUTBOL = ("f5a", "f5b", "f7")
CLAVES_PADEL = ("padelT", "padelA")
ORDEN_CLAVES = {c: i for i, c in enumerate(CLAVES_FUTBOL + CLAVES_PADEL)}
# Sólo si todavía no hay ninguna cancha de ese deporte cargada: ap, ci y duración (minutos) de la planilla de papel.
HORARIOS_PAPEL = {"futbol": (840, 1380, 60), "padel": (600, 1410, 90)}
VERSION_DISENO = 1                                  # subirla si cambia el dibujo de la hoja: así las agendas ya escritas se rehacen
AMARILLO = PatternFill("solid", fgColor="FFFF00")
HORA = re.compile(r"^(\d{1,2})(?::(\d{2}))?\s*hs$", re.IGNORECASE)
PREFIJO_FIRMA = "agenda:"                           # la firma de lo que se escribió queda en las propiedades del archivo: así, al volver a abrir el programa, lo que no cambió no se toca
COLUMNAS_A_MANO = (2, 3, 8, 9)                      # $ y MP de cada cancha (B, C, H, I)


def ventana(hoy: dt.date) -> list[str]:
    """Hoy y los 7 días siguientes."""
    return [(hoy + dt.timedelta(days=i)).isoformat() for i in range(DIAS_ADELANTE + 1)]


# ------------------------------------------------------------------ lectura y armado de los datos

def leer_datos(con: sqlite3.Connection, desde: str, hasta: str) -> dict:
    """Canchas, turnos fijos y los turnos que tocan el rango de fechas (por su fecha o por la fecha de la ocurrencia fija que reemplazan)."""
    def lista(coleccion: str) -> list[dict]:
        return [d for d in (json.loads(r[0]) for r in con.execute("SELECT data FROM docs WHERE coleccion = ? ORDER BY seq", (coleccion,))) if isinstance(d, dict)]

    turnos = [d for d in (json.loads(r[0]) for r in con.execute(
        "SELECT data FROM docs WHERE coleccion = 'turnos' AND ((dia >= ? AND dia <= ?) OR json_extract(data, '$.ocurrencia') BETWEEN ? AND ?) ORDER BY seq",
        (desde, hasta, desde, hasta))) if isinstance(d, dict)]
    return {"canchas": lista("canchas"), "fijos": lista("fijos"), "turnos": turnos}


def _entero(v: Any) -> int | None:
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != v or abs(v) > 10**7:
        return None
    return int(v)


def dow_pantalla(fecha: str) -> int:
    """Día de la semana como lo cuenta la pantalla (JavaScript): domingo = 0 ... sábado = 6."""
    return (dt.date.fromisoformat(fecha).weekday() + 1) % 7


def ocurrencias_fijas(fecha: str, fijos: list[dict], turnos: list[dict], canchas: dict[str, dict]) -> list[dict]:
    """Los turnos fijos que tocan jugar en `fecha`. Es lo mismo que hace `ocurrenciasFijas()` en predio/web/index.html (hay una prueba que
    corre las dos con los mismos datos): un fijo vale ese día si cae en su día de la semana, ya empezó, no terminó y no hay un turno real
    de ese fijo para esa fecha (si se suspendió, se movió o se cobró, el turno real ocupa su lugar)."""
    dow = dow_pantalla(fecha)
    reemplazados = {(t.get("fijoId"), t.get("ocurrencia")) for t in turnos if t.get("fijoId") is not None}
    out = []
    for f in fijos:
        desde, hasta = f.get("desde"), f.get("hasta")
        if f.get("dow") != dow or not isinstance(desde, str) or desde > fecha or (hasta and fecha > hasta):
            continue
        if (f.get("id"), fecha) in reemplazados:
            continue
        c = canchas.get(f.get("canchaId"))
        ini, dur = _entero(f.get("ini")), _entero(c.get("dur")) if c else None
        if c is None or ini is None or dur is None:
            continue
        out.append({"virtual": True, "id": f"F:{f.get('id')}:{fecha}", "fijoId": f.get("id"), "canchaId": f.get("canchaId"), "fecha": fecha, "ini": ini, "fin": ini + dur,
                    "cliente": f.get("cliente"), "tel": f.get("tel"), "nota": f.get("nota"), "esFijo": True, "ocurrencia": fecha})
    return out


def _limpio(v: Any, tope: int = 200) -> str:
    """Texto en una sola línea y de un largo razonable (una celda de Excel no admite más de 32 767 caracteres)."""
    t = " ".join(str(v).split()) if v else ""
    return t if len(t) <= tope else t[:tope - 1] + "…"


def _items_del_dia(dia: str, datos: dict) -> list[dict]:
    """Todo lo que ocupa una cancha ese día (turnos que no están cancelados + los fijos que tocan), ya con la cancha ubicada por su clave."""
    por_id = {c.get("id"): c for c in datos["canchas"]}
    crudos = [t for t in datos["turnos"] if t.get("fecha") == dia and t.get("estado") != "cancelado"]
    crudos += ocurrencias_fijas(dia, datos["fijos"], datos["turnos"], por_id)
    items = []
    for t in crudos:
        c = por_id.get(t.get("canchaId"))
        ini = _entero(t.get("ini"))
        if c is None or ini is None:
            continue
        fin = _entero(t.get("fin"))
        items.append({"clave": c.get("clave") or "", "cancha": _limpio(c.get("nombre")) or "Cancha", "ini": ini, "fin": max(fin if fin is not None else ini + 1, ini + 1),
                      "cliente": _limpio(t.get("cliente"), 120) or "Sin nombre", "nota": _limpio(t.get("nota"), 500), "fijo": bool(t.get("esFijo"))})
    items.sort(key=lambda i: (i["ini"], ORDEN_CLAVES.get(i["clave"], 99), i["cliente"]))
    return items


def _inicios_de_grilla(canchas: list[dict], deporte: str) -> set[int]:
    """Los horarios de turno que tienen las canchas (de su apertura, de a una duración, mientras el turno termine antes del cierre)."""
    inicios: set[int] = set()
    for c in canchas:
        ap, ci, dur = _entero(c.get("ap")), _entero(c.get("ci")), _entero(c.get("dur"))
        if ap is not None and ci is not None and dur and dur > 0:
            inicios.update(range(ap, ci - dur + 1, dur))
    if not canchas:
        ap, ci, dur = HORARIOS_PAPEL[deporte]
        inicios.update(range(ap, ci - dur + 1, dur))
    return inicios


def _celda(items: list[dict]) -> dict | None:
    """Lo que se escribe en una casilla de reserva (si por un dato raro hay dos turnos a la vez, quedan los dos nombres)."""
    if not items:
        return None
    nombres: list[str] = []
    for i in items:
        n = i["cliente"] + (" (F)" if i["fijo"] else "")
        if n not in nombres:
            nombres.append(n)
    return {"t": " / ".join(nombres), "f": any(i["fijo"] for i in items)}


def _marca(valores: list[str]) -> str:
    return "/".join(dict.fromkeys(valores))


def etiqueta_futbol(m: int) -> str:
    h, mm = divmod(m, 60)
    return f"{h} Hs" if mm == 0 else f"{h}:{mm:02d} Hs"


def etiqueta_padel(m: int) -> str:
    return f"{m // 60:02d}:{m % 60:02d} hs"


def armar_dia(dia: str, datos: dict) -> dict:
    """Todo lo que lleva la hoja de un día, como datos simples (sirve para comparar si algo cambió y para dibujarla)."""
    d = dt.date.fromisoformat(dia)
    items = _items_del_dia(dia, datos)
    futbol = [i for i in items if i["clave"] in CLAVES_FUTBOL]
    padel = [i for i in items if i["clave"] in CLAVES_PADEL]
    for_canchas = [c for c in datos["canchas"] if c.get("clave") in CLAVES_FUTBOL]
    pad_canchas = [c for c in datos["canchas"] if c.get("clave") in CLAVES_PADEL]

    filas_f = []
    for ini in sorted(_inicios_de_grilla(for_canchas, "futbol") | {i["ini"] for i in futbol}):
        aca = [i for i in futbol if i["ini"] <= ini < i["fin"]]
        c1 = [i for i in aca if i["clave"] in ("f5a", "f7")]
        c2 = [i for i in aca if i["clave"] in ("f5b", "f7")]
        filas_f.append({"ini": ini, "etiqueta": etiqueta_futbol(ini), "r1": _celda(c1), "r2": _celda(c2),
                        "k1": _marca(["7" if i["clave"] == "f7" else "5" for i in c1]), "k2": _marca(["7" if i["clave"] == "f7" else "5" for i in c2])})
    filas_p = []
    for ini in sorted(_inicios_de_grilla(pad_canchas, "padel") | {i["ini"] for i in padel}):
        aca = [i for i in padel if i["ini"] <= ini < i["fin"]]
        filas_p.append({"ini": ini, "etiqueta": etiqueta_padel(ini), "r1": _celda([i for i in aca if i["clave"] == "padelT"]), "r2": _celda([i for i in aca if i["clave"] == "padelA"])})

    obs = []
    for i in items:
        if i["nota"]:
            obs.append(f"{i['cancha']} · {hm(i['ini'])} hs · {i['nota']}")
        if i["clave"] not in ORDEN_CLAVES:
            obs.append(f"{i['cancha']} · {hm(i['ini'])} hs · {i['cliente']} (esta cancha no tiene lugar en la planilla)")
    return {"v": VERSION_DISENO, "dia": dia, "titulo": f"{DIAS_MAYUS[d.weekday()]}  {d.day:02d} / {d.month:02d} / {d.year}", "futbol": filas_f, "padel": filas_p, "observaciones": obs}


def firma(modelo: dict) -> str:
    return hashlib.sha1(json.dumps(modelo, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()


# ------------------------------------------------------------------ lo que se escribió a mano

def leer_previo(ruta: Path) -> tuple[str | None, dict[tuple[str, int], dict[int, tuple[Any, str]]]]:
    """De una agenda ya escrita: la firma con la que se armó y lo que alguien escribió a mano en las columnas $ y MP, como
    {(bloque, minutos): {columna: (valor, tipo)}}. Si el archivo está bloqueado por otro programa se avisa (ArchivoEnUso) para no pisarlo y
    reintentar; si está dañado o es de otra cosa se ignora (no hay nada que conservar)."""
    try:
        wb = load_workbook(ruta)
    except PermissionError as e:
        raise ArchivoEnUso(f"{ruta.name} está abierto en otro programa") from e
    except Exception:      # noqa: BLE001
        return None, {}
    out: dict[tuple[str, int], dict[int, tuple[Any, str]]] = {}
    try:
        descripcion = wb.properties.description or ""
        previa = descripcion[len(PREFIJO_FIRMA):] if descripcion.startswith(PREFIJO_FIRMA) else None
        ws = wb["Agenda"] if "Agenda" in wb.sheetnames else wb.worksheets[0]
        bloque = "F"
        for fila in ws.iter_rows(min_col=1, max_col=9):
            a = fila[0].value
            if not isinstance(a, str):
                continue
            t = a.strip()
            if t.upper() == "OBSERVACIONES":
                break
            if t.upper() == "PADEL CANCHA 1":
                bloque = "P"
                continue
            m = HORA.match(t)
            if not m:
                continue
            ini = int(m.group(1)) * 60 + int(m.group(2) or 0)
            for col in COLUMNAS_A_MANO:
                c = fila[col - 1]
                if c.value not in (None, "") and c.data_type in ("n", "s", "f", "b", "d"):
                    out.setdefault((bloque, ini), {})[col] = (c.value, c.data_type)
    except Exception:      # noqa: BLE001
        log.warning("No se pudo leer lo escrito a mano en %s", ruta, exc_info=True)
        return None, out
    finally:
        wb.close()
    return previa, out


# ------------------------------------------------------------------ el Excel

_FINO, _MEDIO = Side(style="thin"), Side(style="medium")


def _cuadro(ws, f1: int, c1: int, f2: int, c2: int) -> None:
    """Bordes finos adentro y más gruesos en el contorno (como las tablas del papel)."""
    for f in range(f1, f2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(f, c).border = Border(left=_MEDIO if c == c1 else _FINO, right=_MEDIO if c == c2 else _FINO,
                                          top=_MEDIO if f == f1 else _FINO, bottom=_MEDIO if f == f2 else _FINO)


def _marcos(ws, f1: int, f2: int, columnas: tuple[int, ...]) -> None:
    """Contorno grueso de una casilla (las dos «5_7» van recuadradas en el papel)."""
    for f in range(f1, f2 + 1):
        for c in columnas:
            b = ws.cell(f, c).border
            ws.cell(f, c).border = Border(left=_MEDIO, right=_MEDIO, top=b.top, bottom=b.bottom)


def _poner(ws, f: int, c: int, valor: Any, negrita: bool = False, tam: int = 11, alinear: str = "center", fondo: PatternFill | None = None, achicar: bool = False) -> None:
    cel = ws.cell(f, c, _san(valor) if valor not in (None, "") else None)
    cel.font = Font(name="Calibri", bold=negrita, size=tam)
    cel.alignment = Alignment(horizontal=alinear, vertical="center", shrink_to_fit=achicar)
    if fondo is not None:
        cel.fill = fondo
    if isinstance(cel.value, str) and cel.value.startswith(("=", "+", "-", "@")):
        cel.quotePrefix = True          # si alguien edita la casilla, Excel la sigue tomando por texto


def _reserva(ws, f: int, c: int, celda: dict | None) -> None:
    _poner(ws, f, c, celda["t"] if celda else None, tam=12, alinear="left", fondo=AMARILLO if celda and celda["f"] else None, achicar=True)


def construir_agenda(modelo: dict, destino: Path, manuales: dict | None = None) -> None:
    manuales = manuales or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Agenda"
    for letra, ancho in zip("ABCDEFGHI", (10, 11, 11, 38, 6.5, 6.5, 38, 11, 11)):
        ws.column_dimensions[letra].width = ancho
    ws.sheet_view.showGridLines = False
    a_mano: list[tuple[int, int, tuple[str, int]]] = []          # casillas $ y MP: (fila, columna, (bloque, minutos)); se completan al final con lo que hubiera escrito alguien

    # título
    ws.merge_cells("A1:I1")
    _poner(ws, 1, 1, modelo["titulo"], negrita=True, tam=16)
    ws.row_dimensions[1].height = 30

    # FÚTBOL
    fila = 2
    for c, t in enumerate(("TURNO", "$", "MP", "CANCHA 1", "5_7", "5_7", "CANCHA 2", "$", "MP"), start=1):
        _poner(ws, fila, c, t, negrita=True)
    ws.row_dimensions[fila].height = 22
    primera = fila + 1
    for r in modelo["futbol"]:
        fila += 1
        ws.row_dimensions[fila].height = 21
        _poner(ws, fila, 1, r["etiqueta"], negrita=True)
        _reserva(ws, fila, 4, r["r1"])
        _reserva(ws, fila, 7, r["r2"])
        _poner(ws, fila, 5, r["k1"], negrita=True, tam=12)
        _poner(ws, fila, 6, r["k2"], negrita=True, tam=12)
        for col in COLUMNAS_A_MANO:
            ws.cell(fila, col).alignment = Alignment(horizontal="center", vertical="center")
            if col in (2, 8):
                ws.cell(fila, col).number_format = FORMATO_PLATA
            a_mano.append((fila, col, ("F", r["ini"])))
    _cuadro(ws, 2, 1, fila, 9)
    if fila >= primera:
        _marcos(ws, 2, fila, (5, 6))

    # PÁDEL
    fila += 2
    ws.row_dimensions[fila - 1].height = 12
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    ws.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=9)
    _poner(ws, fila, 1, "PADEL CANCHA 1", negrita=True, tam=13)
    _poner(ws, fila, 7, "PADEL CANCHA 2", negrita=True, tam=13)
    ws.row_dimensions[fila].height = 24
    fila += 1
    cab = fila
    for c, t in ((1, "TURNO"), (2, "$"), (3, "MP"), (4, "RESERVA"), (5, "TURNO"), (7, "RESERVA"), (8, "$"), (9, "MP")):
        _poner(ws, fila, c, t, negrita=True)
    ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
    ws.row_dimensions[fila].height = 22
    for r in modelo["padel"]:
        fila += 1
        ws.row_dimensions[fila].height = 21
        ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
        _poner(ws, fila, 1, r["etiqueta"], negrita=True)
        _poner(ws, fila, 5, r["etiqueta"], negrita=True)
        _reserva(ws, fila, 4, r["r1"])
        _reserva(ws, fila, 7, r["r2"])
        for col in COLUMNAS_A_MANO:
            ws.cell(fila, col).alignment = Alignment(horizontal="center", vertical="center")
            if col in (2, 8):
                ws.cell(fila, col).number_format = FORMATO_PLATA
            a_mano.append((fila, col, ("P", r["ini"])))
    _cuadro(ws, cab, 1, fila, 4)
    _cuadro(ws, cab, 5, fila, 6)
    _cuadro(ws, cab, 7, fila, 9)

    # OBSERVACIONES: una nota por línea, en una caja grande
    fila += 2
    ws.row_dimensions[fila - 1].height = 12
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
    _poner(ws, fila, 1, "OBSERVACIONES", negrita=True, tam=12, alinear="left")
    ws.row_dimensions[fila].height = 22
    ini_obs = fila + 1
    for texto in modelo["observaciones"] + [""] * max(0, 6 - len(modelo["observaciones"])):
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
        _poner(ws, fila, 1, texto, tam=12, alinear="left")
        ws.cell(fila, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[fila].height = 21 * max(1, -(-len(texto) // 120))
    _cuadro(ws, ini_obs, 1, fila, 9)

    # impresión: una hoja A4 apaisada
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:I{fila}"

    _limpiar_hoja(ws)                                           # textos que empiezan con = + - @ quedan como texto, nunca como fórmula
    for f, col, clave in a_mano:                                # lo escrito a mano se vuelve a poner después: ahí una fórmula (=5000+2000) tiene que seguir siendo fórmula
        previo = manuales.get(clave, {}).get(col)
        if previo is not None:
            cel = ws.cell(f, col)
            cel.value = previo[0]
            if previo[1] == "s":
                cel.data_type = "s"                             # lo que era texto sigue siendo texto
    wb.properties.creator = "Caja del Predio"
    wb.properties.title = f"Agenda {modelo['dia']}"
    wb.properties.description = PREFIJO_FIRMA + firma(modelo)
    wb.save(destino)


# ------------------------------------------------------------------ proceso de fondo

class AgendaDiaria:
    """Decide cuándo hay que revisar las agendas y las escribe. ArchivosDiarios la usa desde su mismo hilo de fondo."""

    def __init__(self, diario):
        self.diario = diario
        self._lock = threading.Lock()
        self._trabajo = threading.Lock()              # una sola escritura de agendas a la vez (puede pedirse desde el hilo de fondo y desde la pantalla)
        self._revisar = False
        self._fallidos: set[str] = set()              # días que no se pudieron escribir (archivo abierto): se reintentan
        self._firmas: dict[str, str] = {}             # día -> firma de lo último que se escribió (sólo se toca con _trabajo tomado)
        self._fecha: dt.date | None = None
        self.ultimo_error: str | None = None
        self.ultima_generacion: float | None = None

    # --- pedidos ---
    def pedir(self) -> None:
        """Algo cambió (turnos, fijos, canchas u otra cosa): revisar todas las agendas. Revisar es barato: sólo se escribe la que cambió."""
        with self._lock:
            self._revisar = True

    def cambio_de_dia(self, hoy: dt.date | None = None) -> bool:
        """Verdadero la primera vez y cada vez que cambia la fecha de la PC (arrancó el programa o empezó otro día): hay que armar la ventana de días nueva."""
        hoy = hoy or dt.date.today()
        with self._lock:
            if self._fecha == hoy:
                return False
            self._fecha = hoy
            self._revisar = True
            return True

    def hay_pendiente(self) -> bool:
        """Hay algo que hacer: revisar (pedido nuevo) o reintentar una agenda que estaba abierta."""
        with self._lock:
            return self._revisar or bool(self._fallidos)

    def por_revisar(self) -> bool:
        """Llegó un pedido nuevo que todavía no se atendió (no cuenta lo que quedó trabado por un archivo abierto: eso se reintenta cada tanto)."""
        with self._lock:
            return self._revisar

    def fallidos(self) -> list[str]:
        with self._lock:
            return sorted(self._fallidos)

    # --- escritura ---
    def escribir(self, hoy: dt.date | None = None) -> None:
        with self._trabajo:
            with self._lock:
                revisar, self._revisar = self._revisar, False
                dias = ventana(hoy or dt.date.today())
                self._fallidos &= set(dias)
                self._firmas = {d: f for d, f in self._firmas.items() if d in dias}
                a_ver = dias if revisar else [d for d in dias if d in self._fallidos]
            if not a_ver:
                return
            try:
                con = self.diario.db.lectura()
                try:
                    datos = leer_datos(con, dias[0], dias[-1])
                finally:
                    con.close()
            except Exception as e:      # noqa: BLE001 - un error acá nunca debe frenar la caja
                log.exception("No se pudieron leer los turnos para la agenda")
                with self._lock:
                    self._fallidos |= set(a_ver)           # se reintenta cada tanto (no en el acto: así no se queda dando vueltas)
                    self.ultimo_error = str(e)
                return
            if not datos["canchas"]:                  # base recién creada: todavía no hay canchas, no hay nada que armar
                return
            for dia in a_ver:
                try:
                    self._generar(dia, datos)
                except ArchivoEnUso as e:
                    with self._lock:
                        self._fallidos.add(dia)
                        self.ultimo_error = str(e)
                except Exception as e:      # noqa: BLE001
                    log.exception("No se pudo generar la agenda del %s", dia)
                    with self._lock:
                        self._fallidos.add(dia)
                        self.ultimo_error = str(e)
                else:
                    with self._lock:
                        self._fallidos.discard(dia)
            with self._lock:
                if not self._fallidos:
                    self.ultimo_error = None

    def _generar(self, dia: str, datos: dict) -> bool:
        """Escribe la agenda de un día si cambió desde la última vez. Devuelve si escribió."""
        modelo = armar_dia(dia, datos)
        f = firma(modelo)
        ruta = self.diario.ruta_agenda(dia)
        manuales: dict = {}
        if ruta.exists():
            if self._firmas.get(dia) == f:
                return False
            previa, manuales = leer_previo(ruta)          # (si está bloqueada levanta ArchivoEnUso: no se pisa, se reintenta)
            if previa == f:                              # ya estaba así, por ejemplo de antes de cerrar el programa: no se toca
                self._firmas[dia] = f
                return False
        self.diario.carpeta(dia).mkdir(parents=True, exist_ok=True)

        def escribir(tmp: Path) -> None:
            try:
                construir_agenda(modelo, tmp, manuales)
            except BaseException:
                tmp.unlink(missing_ok=True)
                raise

        _escribir_atomico(ruta, escribir)
        self._firmas[dia] = f
        self.ultima_generacion = time.time()
        return True
