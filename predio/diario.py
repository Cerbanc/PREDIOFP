"""Carpeta de cada día: Excel y registro de cambios, siempre al día y sin que nadie tenga que exportar nada.

    C:\\Predio\\dias\\2026-10-03\\2026-10-03_resumen.xlsx   (resumen, ventas, detalle, caja, stock, turnos, cambios)
    C:\\Predio\\dias\\2026-10-03\\2026-10-03_registro.txt   (todo lo que cambió ese día, en texto simple)

Cada vez que se guarda algo se avisa a este módulo; unos segundos después (para no trabajar de más
cuando hay muchas ventas seguidas) se regeneran los archivos de los días tocados. Si el Excel del día está
abierto en la PC, Windows no deja reemplazarlo: se reintenta solo hasta que se cierre.
"""
from __future__ import annotations

import datetime as dt
import json
import logging
import os
import sqlite3
import threading
import time
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Rutas
from .db import BaseDeDatos
from .util import dia_de, es_dia, hm, hora_de, plata

log = logging.getLogger("predio.diario")

DIAS_SEMANA = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
FORMATO_PLATA = '"$" #,##0;[Red]-"$" #,##0'
VERDE = PatternFill("solid", fgColor="0A7A5A")
GRIS = PatternFill("solid", fgColor="E6ECE9")
NEGRITA_BLANCA = Font(bold=True, color="FFFFFF")


class ArchivoEnUso(OSError):
    pass


def nombre_dia(dia: str) -> str:
    d = dt.date.fromisoformat(dia)
    return f"{DIAS_SEMANA[d.weekday()]} {d.day:02d}/{d.month:02d}/{d.year}"


# ------------------------------------------------------------------ datos del día

def _docs(con: sqlite3.Connection, coleccion: str, dia: str) -> list[dict]:
    return [json.loads(r[0]) for r in con.execute("SELECT data FROM docs WHERE coleccion = ? AND dia = ? ORDER BY ts, seq", (coleccion, dia))]


def _todos(con: sqlite3.Connection, coleccion: str) -> list[dict]:
    return [json.loads(r[0]) for r in con.execute("SELECT data FROM docs WHERE coleccion = ? ORDER BY seq", (coleccion,))]


def _unico(con: sqlite3.Connection, nombre: str, defecto: Any = None) -> Any:
    fila = con.execute("SELECT data FROM docs WHERE coleccion = ? AND id = '_'", (nombre,)).fetchone()
    return json.loads(fila[0]) if fila else defecto


def _limites_dia(dia: str) -> tuple[int, int]:
    d = dt.datetime.strptime(dia, "%Y-%m-%d")
    ini = int(d.timestamp() * 1000)
    return ini, ini + 86_400_000


def datos_del_dia(con: sqlite3.Connection, dia: str) -> dict:
    """Junta todo lo del día y calcula los totales. También lo usa el panel remoto."""
    if not es_dia(dia):
        raise ValueError("Fecha inválida")
    ini_ms, fin_ms = _limites_dia(dia)
    categorias = {c["id"]: c.get("nombre", "") for c in _todos(con, "categorias")}
    ventas = _docs(con, "ventas", dia)
    ledger = _docs(con, "caja", dia)
    mov_caja = _docs(con, "movCaja", dia)
    stock = _docs(con, "movimientos", dia)
    cc = _docs(con, "cc", dia)
    canchas = {c["id"]: c.get("nombre", "") for c in _todos(con, "canchas")}
    turnos = [json.loads(r[0]) for r in con.execute("SELECT data FROM docs WHERE coleccion = 'turnos' AND dia = ? ORDER BY seq", (dia,))]
    cajas = [json.loads(r[0]) for r in con.execute(
        "SELECT data FROM docs WHERE coleccion = 'cajas' AND (dia = ? OR (json_extract(data, '$.cerrada') >= ? AND json_extract(data, '$.cerrada') < ?)) ORDER BY ts",
        (dia, ini_ms, fin_ms))]
    cambios = [dict(r) for r in con.execute(
        "SELECT id, ts, usuario, coleccion, ref, accion, importante, resumen FROM auditoria WHERE dia = ? ORDER BY id", (dia,))]
    productos = _todos(con, "productos")

    ok = [v for v in ventas if v.get("estado") == "ok"]
    anuladas = [v for v in ventas if v.get("estado") != "ok"]
    total = sum(v.get("total") or 0 for v in ok)

    medios: dict[str, dict[str, float]] = defaultdict(lambda: {"cobros": 0, "devoluciones": 0, "correcciones": 0})
    for e in ledger:
        clave = {"cobro": "cobros", "devolucion": "devoluciones", "correccion": "correcciones"}.get(e.get("tipo"), "correcciones")
        medios[e.get("nombre") or "?"][clave] += e.get("monto") or 0
    por_medio = [{"medio": m, **v, "neto": v["cobros"] + v["devoluciones"] + v["correcciones"]} for m, v in sorted(medios.items())]

    cat_tot: dict[str, dict[str, float]] = defaultdict(lambda: {"unidades": 0, "importe": 0})
    prod_tot: dict[str, dict[str, float]] = defaultdict(lambda: {"unidades": 0, "importe": 0})
    for v in ok:
        for it in v.get("items") or []:
            cant, precio = it.get("cant") or 0, it.get("precio") or 0
            c = categorias.get(it.get("catId"), "Sin categoría")
            cat_tot[c]["unidades"] += cant
            cat_tot[c]["importe"] += cant * precio
            prod_tot[it.get("nombre") or "?"]["unidades"] += cant
            prod_tot[it.get("nombre") or "?"]["importe"] += cant * precio
    dest_tot: dict[str, dict[str, float]] = defaultdict(lambda: {"tickets": 0, "total": 0})
    for v in ok:
        d = dest_tot[v.get("destino") or "Mostrador"]
        d["tickets"] += 1
        d["total"] += v.get("total") or 0

    fiado = {"cargos": 0, "pagos": 0, "anulaciones": 0}
    for e in cc:
        k = {"cargo": "cargos", "pago": "pagos", "anulacion": "anulaciones"}.get(e.get("tipo"))
        if k:
            fiado[k] += e.get("monto") or 0

    mov_tot = defaultdict(float)
    for m in mov_caja:
        if not m.get("anulado"):
            mov_tot[m.get("tipo")] += m.get("monto") or 0

    bajos = [{"producto": p.get("nombre"), "stock": p.get("stock"), "minimo": p.get("minimo")} for p in productos
             if p.get("activo") and p.get("tipo") in ("directo", "insumo") and (p.get("stock") or 0) <= (p.get("minimo") or 0)]
    estados = defaultdict(int)
    for t in turnos:
        estados[t.get("estado")] += 1

    return {
        "dia": dia,
        "negocio": _unico(con, "negocio", "Predio"),
        "kpis": {
            "tickets": len(ok), "anulados": len(anuladas), "total_vendido": total,
            "descuentos": sum(v.get("descuento") or 0 for v in ok),
            "ticket_promedio": round(total / len(ok)) if ok else 0,
            "cobrado_neto": sum(m["neto"] for m in por_medio),
            "cambios": len(cambios), "cambios_importantes": sum(1 for c in cambios if c["importante"]),
        },
        "por_medio": por_medio,
        "por_categoria": [{"categoria": k, **v} for k, v in sorted(cat_tot.items(), key=lambda x: -x[1]["importe"])],
        "por_destino": [{"destino": k, **v} for k, v in sorted(dest_tot.items(), key=lambda x: -x[1]["total"])],
        "top_productos": [{"producto": k, **v} for k, v in sorted(prod_tot.items(), key=lambda x: (-x[1]["unidades"], x[0]))[:15]],
        "fiado": fiado,
        "movimientos_caja": dict(mov_tot),
        "turnos_por_estado": dict(estados),
        "stock_bajo": bajos,
        "ventas": ventas, "ledger": ledger, "mov_caja": mov_caja, "stock": stock, "turnos": turnos, "cajas": cajas,
        "cambios": cambios, "canchas": canchas, "categorias": categorias,
    }


# ------------------------------------------------------------------ escritura de archivos

def _escribir_atomico(destino: Path, escribir) -> None:
    tmp = destino.with_name(destino.name + ".tmp")
    escribir(tmp)
    try:
        os.replace(tmp, destino)
    except PermissionError as e:
        tmp.unlink(missing_ok=True)
        raise ArchivoEnUso(f"{destino.name} está abierto en otro programa") from e


def _hoja_tabla(wb: Workbook, nombre: str, columnas: list[str], filas: list[list], plata_cols: Iterable[int] = (), primera: bool = False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = nombre
    ws.append(columnas)
    for c in ws[1]:
        c.font, c.fill, c.alignment = NEGRITA_BLANCA, VERDE, Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(f)
    for i in plata_cols:
        for celdas in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
            for c in celdas:
                c.number_format = FORMATO_PLATA
    for i, col in enumerate(columnas, start=1):
        largo = max([len(str(col))] + [len(str(f[i - 1])) for f in filas[:300] if i - 1 < len(f) and f[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(9, largo + 2))
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _hoja_resumen(wb: Workbook, d: dict, generado: str) -> None:
    ws = wb.active
    ws.title = "Resumen"
    k = d["kpis"]
    fila = [1]

    def linea(*valores, negrita=False, relleno=None, formato=None):
        ws.append(list(valores))
        r = fila[0]
        for c in ws[r]:
            if negrita:
                c.font = Font(bold=True)
            if relleno:
                c.fill = relleno
        if formato:
            for col in formato:
                ws.cell(row=r, column=col).number_format = FORMATO_PLATA
        fila[0] += 1

    def titulo(texto):
        linea()
        linea(texto, negrita=True, relleno=GRIS)

    ws.append([f"Resumen del día · {nombre_dia(d['dia'])}"])
    ws["A1"].font = Font(bold=True, size=15)
    fila[0] = 2
    linea(d["negocio"], f"Actualizado a las {generado}")
    titulo("Ventas")
    linea("Tickets cobrados", k["tickets"])
    linea("Tickets anulados", k["anulados"])
    linea("Total vendido", k["total_vendido"], formato=[2], negrita=True)
    linea("Descuentos dados", k["descuentos"], formato=[2])
    linea("Ticket promedio", k["ticket_promedio"], formato=[2])

    titulo("Plata cobrada por medio de pago (incluye señas y pagos de deuda; no incluye fiado)")
    linea("Medio", "Cobros", "Devoluciones", "Correcciones", "Neto", negrita=True)
    for m in d["por_medio"]:
        linea(m["medio"], m["cobros"], m["devoluciones"], m["correcciones"], m["neto"], formato=[2, 3, 4, 5])
    linea("TOTAL", None, None, None, k["cobrado_neto"], formato=[5], negrita=True)

    titulo("Fiado (cuenta corriente)")
    linea("Fiado generado", d["fiado"]["cargos"], formato=[2])
    linea("Deuda cobrada", d["fiado"]["pagos"], formato=[2])
    linea("Fiado anulado", d["fiado"]["anulaciones"], formato=[2])

    titulo("Caja de efectivo")
    if d["cajas"]:
        linea("Responsable", "Abrió", "Cerró", "Cambio inicial", "Esperado", "Contado", "Diferencia", "Retiro a caja mayor", "Quedó de cambio", negrita=True)
        for c in d["cajas"]:
            linea(c.get("responsable"), hora_de(c["abierta"]) if c.get("abierta") else "", hora_de(c["cerrada"]) if c.get("cerrada") else "abierta",
                  c.get("fondo"), c.get("esperado"), c.get("contado"), c.get("diferencia"), c.get("retiroFinal"), c.get("dejado"), formato=[4, 5, 6, 7, 8, 9])
    else:
        linea("No hubo caja abierta este día")
    mc = d["movimientos_caja"]
    linea("Gastos", mc.get("gasto", 0), formato=[2])
    linea("Pagos a proveedores", mc.get("proveedor", 0), formato=[2])
    linea("Retiros a caja mayor", mc.get("retiro", 0), formato=[2])
    linea("Ingresos de efectivo", mc.get("ingreso", 0), formato=[2])

    titulo("Ventas por categoría (precio de lista, antes de descuentos)")
    linea("Categoría", "Unidades", "Importe", negrita=True)
    for c in d["por_categoria"]:
        linea(c["categoria"], c["unidades"], c["importe"], formato=[3])
    titulo("Ventas por destino")
    linea("Destino", "Tickets", "Total", negrita=True)
    for c in d["por_destino"]:
        linea(c["destino"], c["tickets"], c["total"], formato=[3])
    titulo("Productos más vendidos")
    linea("Producto", "Unidades", "Importe", negrita=True)
    for c in d["top_productos"]:
        linea(c["producto"], c["unidades"], c["importe"], formato=[3])
    titulo("Turnos del día")
    if d["turnos_por_estado"]:
        for estado, n in sorted(d["turnos_por_estado"].items()):
            linea(estado, n)
    else:
        linea("Sin turnos cargados")
    titulo("Stock bajo o agotado (al momento de generar este archivo)")
    if d["stock_bajo"]:
        linea("Producto", "Stock", "Mínimo", negrita=True)
        for b in d["stock_bajo"]:
            linea(b["producto"], b["stock"], b["minimo"])
    else:
        linea("Todo en orden")
    titulo("Cambios de este día")
    linea("Cambios registrados", k["cambios"])
    linea("De los cuales importantes", k["cambios_importantes"])
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 18


def construir_excel(d: dict, destino: Path) -> None:
    wb = Workbook()
    generado = hora_de(time.time() * 1000, True)
    _hoja_resumen(wb, d, generado)
    ventas = sorted(d["ventas"], key=lambda v: v.get("nro") or 0)
    _hoja_tabla(wb, "Ventas", ["Nº", "Hora", "Estado", "Destino", "Cliente", "Subtotal", "Descuento", "Total", "Medios de pago", "Motivo del descuento", "Nota"],
                [[v.get("nro"), hora_de(v["ts"]), v.get("estado"), v.get("destino"), v.get("cliente") or "", v.get("subtotal"), v.get("descuento"), v.get("total"),
                  " + ".join(f"{p.get('nombre')} {plata(p.get('monto'))}" for p in v.get("pagos") or []),
                  ((v.get("ajuste") or {}).get("motivo") or "") + (f" ({v['ajuste']['nota']})" if (v.get("ajuste") or {}).get("nota") else ""), v.get("nota") or ""]
                 for v in ventas], plata_cols=(5, 6, 7))
    filas = []
    for v in ventas:
        for it in v.get("items") or []:
            c, p = it.get("cant") or 0, it.get("precio") or 0
            filas.append([v.get("nro"), hora_de(v["ts"]), v.get("estado"), it.get("nombre"), d["categorias"].get(it.get("catId"), ""), c, p, c * p, it.get("costo") or 0, it.get("nota") or ""])
    _hoja_tabla(wb, "Detalle", ["Ticket", "Hora", "Estado", "Producto", "Categoría", "Cantidad", "Precio", "Subtotal", "Costo unitario", "Nota"], filas, plata_cols=(6, 7, 8))
    tipos = {"cobro": "Cobro", "devolucion": "Devolución", "correccion": "Corrección"}
    _hoja_tabla(wb, "Libro de caja", ["Hora", "Tipo", "Medio de pago", "Monto", "Motivo"],
                [[hora_de(e["ts"]), tipos.get(e.get("tipo"), e.get("tipo")), e.get("nombre"), e.get("monto"), e.get("motivo")] for e in d["ledger"]], plata_cols=(3,))
    nombres_mov = {"gasto": "Gasto", "proveedor": "Proveedor", "retiro": "Retiro a caja mayor", "ingreso": "Ingreso"}
    _hoja_tabla(wb, "Movimientos de caja", ["Hora", "Tipo", "Concepto", "Proveedor", "Monto", "Nota", "Anulado"],
                [[hora_de(m["ts"]), nombres_mov.get(m.get("tipo"), m.get("tipo")), m.get("concepto"), m.get("proveedor") or "", m.get("monto"), m.get("nota") or "", "Sí" if m.get("anulado") else ""]
                 for m in d["mov_caja"]], plata_cols=(4,))
    _hoja_tabla(wb, "Stock", ["Hora", "Producto", "Tipo", "Cantidad", "Antes", "Después", "Detalle"],
                [[hora_de(m["ts"]), m.get("nombre"), m.get("tipo"), m.get("cant"), m.get("antes"), m.get("despues"), m.get("motivo") or ""] for m in d["stock"]])
    _hoja_tabla(wb, "Turnos", ["Hora", "Cancha", "A nombre de", "Precio", "Estado", "Tipo de precio", "Fijo"],
                [[hm(t.get("ini")), d["canchas"].get(t.get("canchaId"), ""), t.get("cliente") or "", t.get("precio"), t.get("estado"), t.get("tipoPrecio") or "", "Sí" if t.get("esFijo") else ""]
                 for t in sorted(d["turnos"], key=lambda t: t.get("ini") or 0)], plata_cols=(3,))
    _hoja_tabla(wb, "Cambios", ["Hora", "Quién", "Importante", "Qué pasó", "Dato", "Acción"],
                [[hora_de(c["ts"], True), c["usuario"], "SÍ" if c["importante"] else "", c["resumen"], c["coleccion"], c["accion"]] for c in d["cambios"]])
    wb.save(destino)


def texto_registro(d: dict) -> str:
    out = [f"Registro de cambios · {nombre_dia(d['dia'])} · {d['negocio']}",
           "Los cambios importantes (borrados, anulaciones, descuentos, plata, precios) empiezan con un !.",
           "Este archivo se regenera solo desde la base de datos; el registro original no se puede borrar.", "-" * 78]
    for c in d["cambios"]:
        out.append(f"{hora_de(c['ts'], True)}  {'!' if c['importante'] else ' '}  {c['usuario']:<14} {c['resumen']}")
    if not d["cambios"]:
        out.append("(sin cambios este día)")
    return "\n".join(out) + "\n"


# ------------------------------------------------------------------ proceso de fondo

class ArchivosDiarios:
    ESPERA = 2.5          # segundos sin novedades antes de escribir
    ESPERA_MAXIMA = 12.0  # aunque sigan llegando cambios, se escribe como mucho cada tanto
    REINTENTO = 30.0

    def __init__(self, db: BaseDeDatos, rutas: Rutas):
        self.db, self.rutas = db, rutas
        self._cond = threading.Condition()
        self._pendientes: dict[str, float] = {}      # día -> no antes de
        self._primera_marca = 0.0
        self._ultima_marca = 0.0
        self._hilo: threading.Thread | None = None
        self._parar = False
        self.ultimo_error: str | None = None
        self.ultima_generacion: float | None = None

    def carpeta(self, dia: str) -> Path:
        return self.rutas.dias / dia

    def ruta_excel(self, dia: str) -> Path:
        return self.carpeta(dia) / f"{dia}_resumen.xlsx"

    def ruta_registro(self, dia: str) -> Path:
        return self.carpeta(dia) / f"{dia}_registro.txt"

    # --- uso ---
    def iniciar(self) -> None:
        self._parar = False
        self._hilo = threading.Thread(target=self._bucle, name="archivos-diarios", daemon=True)
        self._hilo.start()

    def marcar(self, dias: Iterable[str]) -> None:
        ahora = time.monotonic()
        with self._cond:
            for d in dias:
                if es_dia(d):
                    self._pendientes.setdefault(d, 0.0)
            if not self._primera_marca:
                self._primera_marca = ahora
            self._ultima_marca = ahora
            self._cond.notify_all()

    def detener(self, escribir_pendientes: bool = True) -> None:
        with self._cond:
            self._parar = True
            self._cond.notify_all()
        if self._hilo:
            self._hilo.join(timeout=20)
        if escribir_pendientes:
            self.escribir_ahora()

    def pendientes(self) -> list[str]:
        with self._cond:
            return sorted(self._pendientes)

    def escribir_ahora(self) -> None:
        with self._cond:
            dias = list(self._pendientes)
        for d in dias:
            try:
                self.generar_dia(d)
                with self._cond:
                    self._pendientes.pop(d, None)
            except ArchivoEnUso as e:
                self.ultimo_error = str(e)
            except Exception as e:  # noqa: BLE001 - un error acá nunca debe frenar la caja
                self.ultimo_error = str(e)
                log.exception("No se pudo generar el día %s", d)
        with self._cond:
            if not self._pendientes:
                self._primera_marca = self._ultima_marca = 0.0

    # --- generación ---
    def generar_dia(self, dia: str) -> dict:
        con = self.db.lectura()
        try:
            d = datos_del_dia(con, dia)
        finally:
            con.close()
        self.carpeta(dia).mkdir(parents=True, exist_ok=True)
        _escribir_atomico(self.ruta_registro(dia), lambda p: p.write_text(texto_registro(d), encoding="utf-8"))
        _escribir_atomico(self.ruta_excel(dia), lambda p: construir_excel(d, p))
        self.ultimo_error = None
        self.ultima_generacion = time.time()
        return {"excel": str(self.ruta_excel(dia)), "registro": str(self.ruta_registro(dia))}

    def _bucle(self) -> None:
        while True:
            with self._cond:
                while not self._parar:
                    if self._pendientes:
                        ahora = time.monotonic()
                        listo = min(self._ultima_marca + self.ESPERA, self._primera_marca + self.ESPERA_MAXIMA)
                        if ahora >= listo:
                            break
                        self._cond.wait(timeout=max(0.05, listo - ahora))
                    else:
                        self._cond.wait(timeout=5)
                if self._parar:
                    return
            self.escribir_ahora()
            if self.pendientes():                   # quedó algo trabado (archivo abierto): reintentar más tarde
                with self._cond:
                    if not self._parar:
                        self._cond.wait(timeout=self.REINTENTO)
                        self._primera_marca = time.monotonic() - self.ESPERA_MAXIMA
