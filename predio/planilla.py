"""Planilla de productos y stock en Excel: plantilla para completar, exportar lo que hay y leer una planilla completada.

La pantalla "Planilla" muestra la misma tabla y se puede editar ahí mismo; este módulo hace lo mismo con un
archivo .xlsx o .csv para quien prefiera cargar todo en Excel (o en el celular) y después importarlo.
"""
from __future__ import annotations

import csv
import io
import math
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .util import normalizar

# clave interna, título en la planilla, ayuda, ancho
COLUMNAS = [
    ("id", "ID (no tocar)", "Lo completa el sistema. Si dejás un producto sin ID se busca por nombre; si no existe, se crea.", 14),
    ("nombre", "Producto *", "Obligatorio. Ej: Coca-Cola 500 ml", 34),
    ("categoria", "Categoría *", "Elegí de la lista o escribí una nueva (se crea sola). Ej: Kiosco, Pádel, Buffet", 16),
    ("tipo", "Tipo", "Producto (lleva stock propio), Receta (arma con insumos, ej: pancho), Servicio (sin stock, ej: café) o Insumo (no se vende solo, ej: salchicha). Vacío = Producto.", 12),
    ("precio", "Precio de venta", "Lo que paga el cliente por unidad. Los insumos no llevan precio.", 14),
    ("costo", "Costo", "Lo que te cuesta a vos por unidad. Sólo lo ve el administrador. Sirve para vender 'al costo' a empleados.", 12),
    ("stock", "Stock actual", "Unidades que hay HOY. Si el producto ya existe y cambiás este número, queda registrado como ajuste de stock.", 13),
    ("minimo", "Stock mínimo", "El sistema avisa cuando quedan estas unidades o menos.", 13),
    ("paquete", "Cómo viene (paquete)", "Opcional. Ej: Caja x 24, Pack x 6, Bulto 6 x 12", 20),
    ("unidades", "Unidades por paquete", "Cuántas unidades trae el paquete (24, 6, 72). Sirve para cargar remitos por paquete.", 12),
    ("receta", "Receta (sólo tipo Receta)", "Ingredientes que se descuentan al vender. Ej: 1 Salchicha + 1 Pan de pancho. Los nombres tienen que existir como productos o insumos.", 40),
    ("activo", "Activo (SI/NO)", "NO = el producto no aparece en la caja (queda en el historial). Vacío = SI.", 11),
    ("link", "Link (opcional)", "Ficha del producto o del proveedor.", 24),
]
CLAVES = [c[0] for c in COLUMNAS]
TIPOS_TEXTO = {"directo": "Producto", "receta": "Receta", "servicio": "Servicio", "insumo": "Insumo"}

# cómo se reconocen los títulos aunque la planilla venga de otro lado
RECONOCER = [
    ("id", r"^id\b"),
    ("nombre", r"^(producto|nombre|articulo|descripcion)"),
    ("categoria", r"^(categoria|rubro)"),
    ("tipo", r"^tipo"),
    ("precio", r"^precio"),
    ("costo", r"^costo"),
    ("minimo", r"minimo"),
    ("stock", r"^(stock(?!.*min)|cantidad|existencia)"),
    ("unidades", r"(unidades.*(paquete|caja|bulto)|por caja|por bulto|por paquete)"),
    ("paquete", r"^(como viene|paquete|empaque|presentacion)"),
    ("receta", r"^receta"),
    ("activo", r"^(activo|visible)"),
    ("link", r"^(link|url|enlace)"),
]

VERDE = PatternFill("solid", fgColor="0A7A5A")
GRIS = PatternFill("solid", fgColor="E6ECE9")
AMARILLO = PatternFill("solid", fgColor="FFF4CC")


class PlanillaInvalida(ValueError):
    pass


def _tipo_texto(p: dict) -> str:
    return TIPOS_TEXTO.get(p.get("tipo") or "directo", "Producto")


def _receta_texto(p: dict, nombres: dict[str, str]) -> str:
    return " + ".join(f"{r.get('cant', 1)} {nombres.get(r.get('id'), '?')}" for r in (p.get("receta") or []))


def construir_planilla(productos: list[dict], categorias: list[dict], con_datos: bool = True) -> bytes:
    """Arma el .xlsx. Sin datos es la plantilla vacía; con datos es la exportación de lo que hay en el sistema."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append([c[1] for c in COLUMNAS])
    for i, (_, titulo, ayuda, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = VERDE if i > 1 else PatternFill("solid", fgColor="7F8F89")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        celda.comment = Comment(ayuda, "Caja del Predio", width=320, height=110)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"

    nombres = {p["id"]: p.get("nombre", "") for p in productos}
    cat_nombre = {c["id"]: c.get("nombre", "") for c in categorias}
    n = 1
    if con_datos:
        orden = sorted(productos, key=lambda p: (normalizar(cat_nombre.get(p.get("catId"), "")), normalizar(p.get("nombre"))))
        for p in orden:
            conteo = p.get("tipo") in (None, "directo", "insumo")
            ws.append([
                p["id"], p.get("nombre"), cat_nombre.get(p.get("catId"), ""), _tipo_texto(p),
                None if p.get("tipo") == "insumo" else p.get("precio"), p.get("costo") or None,
                p.get("stock") if conteo else None, p.get("minimo") if conteo else None,
                p.get("paqueteNombre") or None, p.get("paqueteUnidades") if (p.get("paqueteUnidades") or 1) > 1 else None,
                _receta_texto(p, nombres) or None, "SI" if p.get("activo", True) else "NO", p.get("link") or None,
            ])
            n += 1
    limite = max(n + 400, 500)
    for fila in ws.iter_rows(min_row=2, max_row=limite):
        for celda in fila:
            if celda.column in (5, 6):
                celda.number_format = '#,##0'
        fila[0].fill = GRIS
    for clave in ("precio", "costo", "stock", "minimo", "unidades"):
        col = get_column_letter(CLAVES.index(clave) + 1)
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            errorTitle="Número", error="Poné un número mayor o igual a cero.")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{limite}")

    listas = wb.create_sheet("Listas")
    listas.append(["Categorías", "Tipos", "Activo"])
    cats = [c.get("nombre", "") for c in categorias] or ["Kiosco", "Pádel", "Buffet"]
    for i in range(max(len(cats), 4)):
        listas.append([cats[i] if i < len(cats) else None, ["Producto", "Receta", "Servicio", "Insumo"][i] if i < 4 else None, ["SI", "NO"][i] if i < 2 else None])
    for c in listas[1]:
        c.font = Font(bold=True)
    listas.column_dimensions["A"].width = 22
    dv_tipo = DataValidation(type="list", formula1="=Listas!$B$2:$B$5", allow_blank=True)
    dv_act = DataValidation(type="list", formula1="=Listas!$C$2:$C$3", allow_blank=True)
    dv_cat = DataValidation(type="list", formula1=f"=Listas!$A$2:$A${max(len(cats), 4) + 1}", allow_blank=True, showErrorMessage=False)
    for dv, clave in ((dv_tipo, "tipo"), (dv_act, "activo"), (dv_cat, "categoria")):
        ws.add_data_validation(dv)
        col = get_column_letter(CLAVES.index(clave) + 1)
        dv.add(f"{col}2:{col}{limite}")

    ins = wb.create_sheet("Instrucciones", 0)
    renglones = [
        ("Planilla de productos y stock", True),
        ("", False),
        ("1. Completá la hoja «Productos»: una fila por producto. Sólo son obligatorios el nombre y la categoría.", False),
        ("2. Guardá el archivo y en el programa entrá a Planilla → «Cargar archivo de Excel». Se muestra todo para revisar antes de guardar nada.", False),
        ("3. Si el producto ya existe se actualiza (si cambia el stock queda registrado como ajuste). Si no existe se crea con ese stock como stock inicial.", False),
        ("", False),
        ("Tipos de producto", True),
        ("Producto: lleva su propio stock (gaseosas, golosinas, pelotas).", False),
        ("Receta: se arma con insumos (pancho = 1 salchicha + 1 pan). Se descuentan los insumos al vender. Se completa la columna Receta.", False),
        ("Servicio: no lleva stock (café, minutas, alquiler de paleta).", False),
        ("Insumo: sólo se usa dentro de recetas (salchichas, panes). No tiene precio de venta.", False),
        ("", False),
        ("Consejos", True),
        ("• La columna ID la completa el sistema al exportar. No la cambies: sirve para no duplicar productos si cambiás un nombre.", False),
        ("• También podés copiar y pegar celdas desde Excel directamente en la pantalla Planilla del programa.", False),
        ("• Los precios van sin $ ni puntos (2200). El stock en unidades sueltas (una caja de 24 son 24).", False),
        ("• Para ocultar un producto sin perder su historial poné NO en la columna Activo.", False),
    ]
    for texto, negrita in renglones:
        ins.append([texto])
        if negrita:
            ins.cell(row=ins.max_row, column=1).font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 130
    for fila in ins.iter_rows():
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ lectura

def numero_ar(v: Any) -> int | None:
    """Convierte lo que venga en una celda a un número entero de pesos/unidades. Vacío = None. Los negativos se conservan: la pantalla los marca como error."""
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return math.floor(v + 0.5)
    t = re.sub(r"[^\d.,-]", "", str(v))
    if t in ("", "-", ".", ","):
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")
    try:
        return math.floor(float(t) + 0.5)       # como Math.round de la pantalla: 2200,5 -> 2201
    except ValueError:
        return None


def _texto(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _filas_de_archivo(contenido: bytes, nombre: str) -> list[list[Any]]:
    nombre = (nombre or "").lower()
    if nombre.endswith((".csv", ".txt", ".tsv")):
        for cod in ("utf-8-sig", "cp1252"):
            try:
                texto = contenido.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise PlanillaInvalida("No pude leer el texto del archivo")
        primera = texto.splitlines()[0] if texto.strip() else ""
        sep = "\t" if "\t" in primera else (";" if primera.count(";") >= primera.count(",") else ",")
        return [list(r) for r in csv.reader(io.StringIO(texto), delimiter=sep)]
    if not nombre.endswith((".xlsx", ".xlsm")):
        raise PlanillaInvalida("El archivo tiene que ser de Excel (.xlsx) o CSV")
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl lanza varios tipos
        raise PlanillaInvalida("No pude abrir el archivo de Excel. ¿Está dañado o protegido con contraseña?") from e
    hoja = next((wb[n] for n in wb.sheetnames if normalizar(n) == "productos"), None) or wb[wb.sheetnames[0]]
    return [list(fila) for fila in hoja.iter_rows(values_only=True)]


def leer_planilla(contenido: bytes, nombre: str) -> dict:
    """Devuelve las filas de la planilla ya interpretadas (sin tocar la base: la pantalla las revisa y las guarda)."""
    filas = _filas_de_archivo(contenido, nombre)
    cab_idx = None
    for i, f in enumerate(filas[:15]):
        celdas = [normalizar(_texto(c)) for c in f]
        if any(re.match(RECONOCER[1][1], c) for c in celdas):
            cab_idx = i
            break
    if cab_idx is None:
        raise PlanillaInvalida("No encuentro la columna «Producto» (o «Nombre») en los títulos")
    cab = [normalizar(_texto(c)) for c in filas[cab_idx]]
    cols: dict[str, int] = {}
    for clave, patron in RECONOCER:
        for i, h in enumerate(cab):
            if i in cols.values():
                continue
            if h and re.search(patron, h):
                cols[clave] = i
                break
    out = []
    for n, f in enumerate(filas[cab_idx + 1:], start=cab_idx + 2):
        if not any(_texto(c) for c in f):
            continue

        def g(clave):
            i = cols.get(clave)
            return f[i] if i is not None and i < len(f) else None

        fila: dict[str, Any] = {"n": n}
        if "id" in cols:
            fila["id"] = _texto(g("id")) or None
        fila["nombre"] = _texto(g("nombre"))
        for clave in ("categoria", "paquete", "receta", "link"):
            if clave in cols:
                fila[clave] = _texto(g(clave))
        if "tipo" in cols:
            t = normalizar(_texto(g("tipo")))
            fila["tipo"] = "receta" if t.startswith("rec") else "servicio" if t.startswith("serv") else "insumo" if t.startswith("ins") else "directo"
        for clave in ("precio", "costo", "stock", "minimo", "unidades"):
            if clave in cols:
                fila[clave] = numero_ar(g(clave))
        if "activo" in cols:
            a = normalizar(_texto(g("activo")))
            fila["activo"] = not (a in ("no", "n", "0", "false", "falso", "oculto", "inactivo"))
        out.append(fila)
    if not out:
        raise PlanillaInvalida("La planilla no tiene productos debajo de los títulos")
    return {"filas": out, "columnas": ["nombre"] + [c for c in CLAVES if c in cols and c != "nombre"]}
