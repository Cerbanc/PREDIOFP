"""Agenda de turnos en Excel con el diseño de la planilla de papel (predio/agenda.py) y los horarios reales de las canchas."""
import datetime as dt
import json
import re
import shutil
import sqlite3
import subprocess
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import base_inicial, op
from predio import agenda as ag
from predio.diario import ArchivosDiarios
from test_servidor import pedir, srv  # noqa: F401  (srv es una prueba-fixture de test_servidor)

RAIZ = Path(__file__).resolve().parents[1]
HOY = dt.date(2026, 10, 15)          # jueves
F = HOY.isoformat()
PADEL_HORAS = ["10:00 hs", "11:30 hs", "13:00 hs", "14:30 hs", "16:00 hs", "17:30 hs", "19:00 hs", "20:30 hs", "22:00 hs"]
CANCHAS = [("cn_f5a", "f5a", "Fútbol 5 · Cancha 1", "futbol", 60, 840, 1380), ("cn_f5b", "f5b", "Fútbol 5 · Cancha 2", "futbol", 60, 840, 1380),
           ("cn_f7", "f7", "Fútbol 7 · Cancha completa", "futbol", 60, 840, 1380), ("cn_pT", "padelT", "Pádel · Cancha 1", "padel", 90, 600, 1410),
           ("cn_pA", "padelA", "Pádel · Cancha 2", "padel", 90, 600, 1410)]
DUR = {c[0]: c[4] for c in CANCHAS}


def cancha(id_, clave, nombre, deporte, dur, ap, ci):
    return op("canchas", id_, {"id": id_, "clave": clave, "nombre": nombre, "deporte": deporte, "dur": dur, "ap": ap, "ci": ci, "precio": 28000, "precioFijo": 24000, "lugarId": None})


def turno(id_, canchaId, ini, cliente="X", fecha=F, estado="reservado", nota="", fijo=False, fijoId=None, ocurrencia=None, fin=None):
    return op("turnos", id_, {"id": id_, "canchaId": canchaId, "fecha": fecha, "ini": ini, "fin": fin if fin is not None else ini + DUR[canchaId], "cliente": cliente, "tel": "",
                              "nota": nota, "precio": 28000, "estado": estado, "pagos": [], "ventaId": None, "esFijo": fijo, "fijoId": fijoId, "ocurrencia": ocurrencia,
                              "clienteId": None, "tipoPrecio": "fijo" if fijo else "comun"})


def fijo(id_, canchaId, dow, ini, cliente="Fijo", desde="2026-01-01", hasta=None, nota=""):
    return op("fijos", id_, {"id": id_, "canchaId": canchaId, "dow": dow, "ini": ini, "cliente": cliente, "tel": "", "nota": nota, "desde": desde, "hasta": hasta, "clienteId": None})


def montar(almacen, *extra, canchas=CANCHAS):
    almacen.aplicar(base_inicial(), "Caja", 0)
    almacen.aplicar({"ops": [cancha(*c) for c in canchas] + list(extra)}, "Caja", None)
    return ArchivosDiarios(almacen.db, almacen.rutas)


def hoja(ad, dia=F, hoy=HOY):
    """Arma las agendas (como si hoy fuera `hoy`) y devuelve la hoja del día pedido."""
    ad.agenda.pedir()
    ad.agenda.escribir(hoy)
    return load_workbook(ad.ruta_agenda(dia))["Agenda"]


def mapa(ws):
    """Fila de cada horario: {'F': {'14 Hs': 3, ...}, 'P': {'10:00 hs': 17, ...}}"""
    out: dict = {"F": {}, "P": {}}
    bloque = "F"
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == "PADEL CANCHA 1":
            bloque = "P"
        elif a == "OBSERVACIONES":
            break
        elif isinstance(a, str) and re.match(r"^\d", a):
            out[bloque][a] = r
    return out


def reservas(ws, fila):
    """D (reserva cancha 1), E y F (las casillas 5_7), G (reserva cancha 2)."""
    return [ws.cell(fila, c).value for c in (4, 5, 6, 7)]


def amarillo(cel):
    return cel.fill.fill_type == "solid" and cel.fill.fgColor.rgb[-6:] == "FFFF00"


def observaciones(ws):
    r = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "OBSERVACIONES")
    return [ws.cell(i, 1).value for i in range(r + 1, ws.max_row + 1) if ws.cell(i, 1).value]


# ------------------------------------------------------------------ el diseño

def test_dia_vacio_tiene_el_diseno_de_la_planilla_de_papel(almacen, rutas):
    ad = montar(almacen)
    ws = hoja(ad)
    assert ws.title == "Agenda" and ws["A1"].value == "JUEVES  15 / 10 / 2026"
    assert [ws.cell(2, c).value for c in range(1, 10)] == ["TURNO", "$", "MP", "CANCHA 1", "5_7", "5_7", "CANCHA 2", "$", "MP"]
    m = mapa(ws)
    assert list(m["F"]) == [f"{h} Hs" for h in range(14, 23)]
    assert list(m["P"]) == PADEL_HORAS
    for fila in m["F"].values():                      # todo vacío menos el horario
        assert all(ws.cell(fila, c).value is None for c in range(2, 10))
    for fila in m["P"].values():                      # la columna central de pádel repite el horario y todo lo demás está vacío
        assert ws.cell(fila, 5).value == ws.cell(fila, 1).value and all(ws.cell(fila, c).value is None for c in (2, 3, 4, 6, 7, 8, 9))
    pt = next(r for r in range(1, ws.max_row) if ws.cell(r, 1).value == "PADEL CANCHA 1")
    assert ws.cell(pt, 7).value == "PADEL CANCHA 2"
    assert [ws.cell(pt + 1, c).value for c in range(1, 10)] == ["TURNO", "$", "MP", "RESERVA", "TURNO", None, "RESERVA", "$", "MP"]
    assert observaciones(ws) == []
    # una sola hoja A4 apaisada, ajustada a una página
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9
    assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 1
    assert ws.print_area == f"'Agenda'!$A$1:$I${ws.max_row}"
    # celdas combinadas del título, de las dos canchas de pádel, de la columna central y de la caja de observaciones
    combinadas = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:I1", f"A{pt}:D{pt}", f"G{pt}:I{pt}", f"E{pt + 1}:F{pt + 1}"} <= combinadas and all(f"E{f}:F{f}" in combinadas for f in m["P"].values())
    # bordes: finos adentro, gruesos en el contorno y en las dos casillas 5_7
    assert ws["A2"].border.left.style == "medium" and ws["I2"].border.right.style == "medium" and ws["D3"].border.left.style == "thin"
    assert ws["E3"].border.left.style == "medium" and ws["F3"].border.right.style == "medium"
    assert ws["A2"].border.top.style == "medium" and ws.cell(max(m["F"].values()), 1).border.bottom.style == "medium"
    # archivo: dias/AAAA-MM-DD/AAAA-MM-DD_agenda.xlsx, sin temporales
    assert ad.ruta_agenda(F) == rutas.dias / F / f"{F}_agenda.xlsx" and ad.ruta_agenda(F).exists()
    assert not list(rutas.dias.rglob("*.tmp"))


def test_horarios_salen_de_la_configuracion_de_las_canchas(almacen):
    """Las filas no están fijas en el código: si las canchas abren y cierran a otra hora, la agenda cambia."""
    otras = [("cn_f5a", "f5a", "Fútbol 5 · Cancha 1", "futbol", 60, 900, 1260), ("cn_f5b", "f5b", "Fútbol 5 · Cancha 2", "futbol", 60, 900, 1260),
             ("cn_f7", "f7", "Fútbol 7 · Cancha completa", "futbol", 60, 900, 1260), ("cn_pT", "padelT", "Pádel · Cancha 1", "padel", 60, 600, 720),
             ("cn_pA", "padelA", "Pádel · Cancha 2", "padel", 60, 600, 720)]
    ws = hoja(montar(almacen, canchas=otras))
    m = mapa(ws)
    assert list(m["F"]) == ["15 Hs", "16 Hs", "17 Hs", "18 Hs", "19 Hs", "20 Hs"] and list(m["P"]) == ["10:00 hs", "11:00 hs"]


def test_si_faltan_las_canchas_de_padel_sus_filas_son_las_de_la_planilla_de_papel(almacen):
    ws = hoja(montar(almacen, canchas=CANCHAS[:3]))
    assert list(mapa(ws)["P"]) == PADEL_HORAS and list(mapa(ws)["F"]) == [f"{h} Hs" for h in range(14, 23)]


def test_si_faltan_las_canchas_de_futbol_sus_filas_son_las_de_la_planilla_de_papel(almacen):
    ws = hoja(montar(almacen, canchas=CANCHAS[3:]))
    assert list(mapa(ws)["P"]) == PADEL_HORAS and list(mapa(ws)["F"]) == [f"{h} Hs" for h in range(14, 23)]


def test_textos_larguisimos_no_rompen_el_archivo(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 1140, "N" * 5000, nota="x" * 50_000)))
    nombre = reservas(ws, mapa(ws)["F"]["19 Hs"])[0]
    assert len(nombre) == 120 and nombre.endswith("…")
    assert all(len(str(c.value)) < 600 for fila in ws.iter_rows() for c in fila if c.value is not None)


def test_futbol_5_en_una_cancha(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno Hernan")))
    f = mapa(ws)["F"]["19 Hs"]
    assert reservas(ws, f) == ["Bruno Hernan", "5", None, None]
    assert not amarillo(ws.cell(f, 4))
    assert reservas(ws, mapa(ws)["F"]["18 Hs"]) == [None, None, None, None]


def test_futbol_5_en_las_dos_canchas(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 1080, "Bruno Hernan"), turno("t2", "cn_f5b", 1080, "Roman Charez")))
    assert reservas(ws, mapa(ws)["F"]["18 Hs"]) == ["Bruno Hernan", "5", "5", "Roman Charez"]


def test_futbol_5_solo_en_la_cancha_2(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5b", 1080, "Roman Charez")))
    assert reservas(ws, mapa(ws)["F"]["18 Hs"]) == [None, None, "5", "Roman Charez"]


def test_futbol_7_repite_el_nombre_y_el_7_en_las_dos_casillas(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f7", 1140, "Guillermo Hernandez")))
    m = mapa(ws)["F"]
    assert reservas(ws, m["19 Hs"]) == ["Guillermo Hernandez", "7", "7", "Guillermo Hernandez"]
    assert reservas(ws, m["20 Hs"]) == [None, None, None, None] and reservas(ws, m["18 Hs"]) == [None, None, None, None]


def test_las_columnas_de_plata_y_mp_quedan_vacias(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f7", 1140, "Guillermo", estado="cobrado"), turno("t2", "cn_pT", 1050, "Ana")))
    m = mapa(ws)
    for fila in list(m["F"].values()) + list(m["P"].values()):
        assert [ws.cell(fila, c).value for c in (2, 3, 8, 9)] == [None, None, None, None]


def test_turno_fijo_lleva_F_y_celda_amarilla(almacen):
    ws = hoja(montar(almacen,
                     fijo("fx1", "cn_f5a", 4, 1200, "Ariel Rimao"),                                                    # jueves 20:00 (la fecha es jueves)
                     turno("t1", "cn_f5b", 1260, "Tobias Vasallo", fijo=True, fijoId="fx9", ocurrencia=F),             # fijo ya cobrado o tocado: es un turno real marcado como fijo
                     fijo("fx3", "cn_f7", 4, 1320, "Mauricio"),
                     turno("t2", "cn_f5a", 1140, "Comun")))
    m = mapa(ws)["F"]
    assert reservas(ws, m["20 Hs"]) == ["Ariel Rimao (F)", "5", None, None] and amarillo(ws.cell(m["20 Hs"], 4)) and not amarillo(ws.cell(m["20 Hs"], 7))
    assert reservas(ws, m["21 Hs"]) == [None, None, "5", "Tobias Vasallo (F)"] and amarillo(ws.cell(m["21 Hs"], 7))
    assert reservas(ws, m["22 Hs"]) == ["Mauricio (F)", "7", "7", "Mauricio (F)"] and amarillo(ws.cell(m["22 Hs"], 4)) and amarillo(ws.cell(m["22 Hs"], 7))
    assert reservas(ws, m["19 Hs"])[0] == "Comun" and not amarillo(ws.cell(m["19 Hs"], 4))      # lo común no se resalta


def test_turno_fijo_en_padel(almacen):
    ws = hoja(montar(almacen, fijo("fx1", "cn_pT", 4, 1050, "Mauricio Nasivera"), fijo("fx2", "cn_pA", 4, 1140, "Lucas")))
    m = mapa(ws)["P"]
    assert ws.cell(m["17:30 hs"], 4).value == "Mauricio Nasivera (F)" and amarillo(ws.cell(m["17:30 hs"], 4))
    assert ws.cell(m["19:00 hs"], 7).value == "Lucas (F)" and amarillo(ws.cell(m["19:00 hs"], 7))


def test_fijo_suspendido_movido_terminado_o_futuro_no_aparece(almacen):
    ad = montar(almacen,
                     fijo("fx1", "cn_f5a", 4, 840, "Suspendido"),
                     turno("t1", "cn_f5a", 840, "Suspendido", estado="cancelado", fijo=True, fijoId="fx1", ocurrencia=F),            # se suspendió ese día
                     fijo("fx2", "cn_f5a", 4, 900, "Movido"),
                     turno("t2", "cn_f5a", 1260, "Movido", fecha="2026-10-16", fijo=True, fijoId="fx2", ocurrencia=F),               # se pasó a otro día
                     fijo("fx3", "cn_f5a", 4, 960, "Terminado", hasta="2026-10-14"),
                     fijo("fx4", "cn_f5a", 4, 1020, "Futuro", desde="2026-10-22"),
                     fijo("fx5", "cn_f5a", 3, 1080, "Otro dia"),
                     fijo("fx6", "cn_f5a", 4, 1140, "Termina hoy", hasta=F),
                     fijo("fx7", "cn_f5a", 4, 1200, "Sin fecha de fin", hasta=""))
    ws = hoja(ad)
    m = mapa(ws)["F"]
    nombres = [ws.cell(f, 4).value for f in m.values()]
    assert [n for n in nombres if n] == ["Termina hoy (F)", "Sin fecha de fin (F)"]
    # el que se movió aparece el día nuevo, y el fijo que empieza más adelante aparece cuando llega su fecha
    ws2 = hoja(ad, "2026-10-16")
    assert reservas(ws2, mapa(ws2)["F"]["21 Hs"])[0] == "Movido (F)"
    ws3 = hoja(ad, "2026-10-22")
    assert reservas(ws3, mapa(ws3)["F"]["17 Hs"])[0] == "Futuro (F)"


def test_el_turno_reprogramado_fuera_de_la_ventana_igual_suspende_la_ocurrencia(almacen):
    ws = hoja(montar(almacen, fijo("fx1", "cn_f5a", 4, 1200, "Se mudó"),
                     turno("t1", "cn_f5a", 1200, "Se mudó", fecha="2027-03-04", fijo=True, fijoId="fx1", ocurrencia=F)))
    assert reservas(ws, mapa(ws)["F"]["20 Hs"]) == [None, None, None, None]


def test_turno_cancelado_no_aparece_y_cobrado_si(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 1140, "Cancelado", estado="cancelado"), turno("t2", "cn_f5b", 1200, "Cobrado", estado="cobrado"),
                     turno("t3", "cn_f5a", 1260, "Otro dia", fecha="2026-10-16")))
    m = mapa(ws)["F"]
    assert reservas(ws, m["19 Hs"]) == [None] * 4 and reservas(ws, m["20 Hs"]) == [None, None, "5", "Cobrado"] and reservas(ws, m["21 Hs"]) == [None] * 4


def test_padel_de_90_minutos(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_pT", 1050, "Mauricio"), turno("t2", "cn_pA", 1320, "Karen"), turno("t3", "cn_pT", 600, "Temprano"), turno("t4", "cn_pA", 690, "Karen 2")))
    m = mapa(ws)["P"]
    assert list(m) == PADEL_HORAS                       # el último turno (22:00 a 23:30) tiene su fila y no se agrega ninguna más
    assert [(h, ws.cell(f, 4).value, ws.cell(f, 7).value) for h, f in m.items() if ws.cell(f, 4).value or ws.cell(f, 7).value] == [
        ("10:00 hs", "Temprano", None), ("11:30 hs", None, "Karen 2"), ("17:30 hs", "Mauricio", None), ("22:00 hs", None, "Karen")]
    # el turno de 17:30 a 19:00 no ocupa la fila de las 19:00 (el fin no cuenta)
    assert ws.cell(m["19:00 hs"], 4).value is None


def test_observaciones_una_nota_por_linea(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno", nota="Trae pelota propia"), turno("t2", "cn_pA", 1140, "Alejandro", nota="Cumpleaños"),
                     turno("t3", "cn_f5b", 1200, "Sin nota"), fijo("fx1", "cn_pT", 4, 1050, "Mauricio", nota="Paga por transferencia"),
                     turno("t4", "cn_f7", 1320, "Cancelado con nota", nota="No va", estado="cancelado"), turno("t5", "cn_f5a", 840, "Otro dia", nota="No es hoy", fecha="2026-10-16")))
    assert observaciones(ws) == ["Pádel · Cancha 1 · 17:30 hs · Paga por transferencia", "Fútbol 5 · Cancha 1 · 19:00 hs · Trae pelota propia", "Pádel · Cancha 2 · 19:00 hs · Cumpleaños"]
    assert ws.cell(next(r for r in range(1, ws.max_row) if ws.cell(r, 1).value == "OBSERVACIONES") + 1, 1).border.left.style == "medium"      # caja con borde


def test_la_caja_de_observaciones_crece_con_las_notas(almacen):
    notas = [turno(f"t{i}", "cn_f5a" if i % 2 else "cn_f5b", 840 + 60 * (i // 2), f"Cliente {i}", nota=f"nota número {i}") for i in range(2, 14)]
    ws = hoja(montar(almacen, *notas))
    assert len(observaciones(ws)) == 12 and observaciones(ws)[0] == "Fútbol 5 · Cancha 1 · 15:00 hs · nota número 3"
    assert ws.max_row == 31 + (12 - 6)                # la caja vacía tiene 6 líneas (hoja de 31 filas): con 12 notas crece a 12


def test_turno_fuera_de_horario_agrega_su_fila(almacen):
    ws = hoja(montar(almacen, turno("t1", "cn_f5a", 720, "Madrugador"), turno("t2", "cn_f5b", 870, "A las dos y media"), turno("t3", "cn_pT", 510, "Temprano"),
                     turno("t4", "cn_pA", 1380, "Tarde")))
    m = mapa(ws)
    assert list(m["F"]) == ["12 Hs"] + [f"{h} Hs" for h in range(14, 15)] + ["14:30 Hs"] + [f"{h} Hs" for h in range(15, 23)]
    assert reservas(ws, m["F"]["12 Hs"]) == ["Madrugador", "5", None, None] and reservas(ws, m["F"]["14:30 Hs"]) == [None, None, "5", "A las dos y media"]
    assert reservas(ws, m["F"]["15 Hs"]) == [None, None, "5", "A las dos y media"]      # y también la fila de las 15:00, que cae dentro de 14:30 a 15:30
    assert list(m["P"]) == ["08:30 hs"] + PADEL_HORAS + ["23:00 hs"]
    assert ws.cell(m["P"]["08:30 hs"], 4).value == "Temprano" and ws.cell(m["P"]["23:00 hs"], 7).value == "Tarde"
    # las filas siguen en orden de horario (la hoja se lee de arriba hacia abajo)
    assert list(m["F"].values()) == sorted(m["F"].values()) and list(m["P"].values()) == sorted(m["P"].values())


def test_un_turno_ocupa_todas_las_filas_que_toca(almacen):
    """Una fila está ocupada si su horario cae dentro de [inicio, fin): un turno de dos horas llena dos filas y el fin no cuenta."""
    ws = hoja(montar(almacen, turno("t1", "cn_f7", 1140, "Torneo", fin=1260)))
    m = mapa(ws)["F"]
    assert [reservas(ws, m[h])[0] for h in ("18 Hs", "19 Hs", "20 Hs", "21 Hs")] == [None, "Torneo", "Torneo", None]


def test_nombres_con_formulas_quedan_como_texto(almacen, rutas):
    malos = ['=HYPERLINK("http://evil.example","clic")', "@SUM(1+1)", "+cmd|' /C calc'!A0", "-2+3"]
    extra = [turno(f"t{i}", "cn_f5a", 840 + 60 * i, n, nota="=1+1" if i == 0 else "") for i, n in enumerate(malos)]
    extra.append(fijo("fx1", "cn_pT", 4, 1050, "=2+2"))
    extra.append(turno("t9", "cn_f5b", 1380, "con\x0bcaracter ilegal"))
    extra.append(turno("t10", "cn_f5b", 1440, "sin\x00nulo"))
    ad = montar(almacen, *extra)
    ws = hoja(ad)
    celdas = [c for fila in ws.iter_rows() for c in fila]
    assert not [c for c in celdas if c.data_type == "f"]
    valores = {c.value for c in celdas if isinstance(c.value, str)}
    assert set(malos) <= valores and "=2+2 (F)" in valores and "con caracter ilegal" in valores and "sinnulo" in valores
    assert all(c.quotePrefix for c in celdas if isinstance(c.value, str) and c.value.startswith(tuple("=+-@")))
    xml = zipfile.ZipFile(ad.ruta_agenda(F)).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in xml and "<f " not in xml and "\x0b" not in xml
    assert "Fútbol 5 · Cancha 1 · 14:00 hs · =1+1" in observaciones(ws)


# ------------------------------------------------------------------ cuándo se arma

def test_se_arman_hoy_y_los_7_dias_siguientes(almacen, rutas):
    ad = montar(almacen)
    hoja(ad)
    existentes = sorted(p.parent.name for p in rutas.dias.glob("*/*_agenda.xlsx"))
    assert existentes == [(HOY + dt.timedelta(days=i)).isoformat() for i in range(8)] and F in existentes and "2026-10-22" in existentes
    assert "2026-10-14" not in existentes and "2026-10-23" not in existentes
    ws = load_workbook(ad.ruta_agenda("2026-10-17"))["Agenda"]
    assert ws["A1"].value == "SÁBADO  17 / 10 / 2026"


def test_sin_canchas_todavia_no_se_arma_nada(almacen, rutas):
    almacen.aplicar(base_inicial(), "Caja", 0)
    ad = ArchivosDiarios(almacen.db, rutas)
    ad.agenda.pedir()
    ad.agenda.escribir(HOY)
    assert not list(rutas.dias.glob("*/*_agenda.xlsx")) and not ad.agenda.hay_pendiente()


def contar_construcciones(monkeypatch):
    """Anota de qué días se rehizo la agenda."""
    llamadas = []
    real = ag.construir_agenda

    def contando(modelo, destino, manuales=None):
        llamadas.append(modelo["dia"])
        real(modelo, destino, manuales)

    monkeypatch.setattr(ag, "construir_agenda", contando)
    return llamadas


def celda_reserva(ad, dia, etiqueta, columna=4):
    ws = load_workbook(ad.ruta_agenda(dia))["Agenda"]
    return ws.cell(mapa(ws)["F"][etiqueta], columna).value


def test_se_regenera_cuando_cambia_un_turno_un_fijo_o_una_cancha_y_solo_ese_dia(almacen, rutas, monkeypatch):
    hoy = dt.date.today()
    dia, manana, en_tres = hoy.isoformat(), (hoy + dt.timedelta(days=1)).isoformat(), (hoy + dt.timedelta(days=3)).isoformat()
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Primero", fecha=dia), turno("t2", "cn_f5a", 1140, "Otro dia", fecha=en_tres))
    llamadas = contar_construcciones(monkeypatch)
    ad.marcar([dia])
    ad.escribir_ahora()
    assert sorted(llamadas) == sorted(ag.ventana(hoy))                          # la primera vez se arman las 8
    assert celda_reserva(ad, dia, "19 Hs") == "Primero"
    llamadas.clear()
    ad.marcar([dia])                                                            # un guardado cualquiera (una venta, por ejemplo): nada cambió en las agendas, no se toca ningún archivo
    ad.escribir_ahora()
    assert llamadas == []
    # un turno cambia de nombre: se rehace sólo la agenda de ese día
    res = almacen.aplicar({"ops": [turno("t1", "cn_f5a", 1140, "Segundo", fecha=dia)]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.escribir_ahora()
    assert llamadas == [dia]
    assert celda_reserva(ad, dia, "19 Hs") == "Segundo"
    # un turno se pasa de un día a otro: se rehacen los dos
    llamadas.clear()
    res = almacen.aplicar({"ops": [turno("t2", "cn_f5a", 1140, "Otro dia", fecha=manana)]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.escribir_ahora()
    assert sorted(llamadas) == sorted([manana, en_tres])
    # un turno fijo nuevo (los fijos no tienen día propio): se revisan todas y se rehacen las que tocan
    llamadas.clear()
    dow = (hoy.weekday() + 1) % 7
    res = almacen.aplicar({"ops": [fijo("fx1", "cn_f5b", dow, 1200, "Fijo nuevo", desde=hoy.isoformat())]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.escribir_ahora()
    assert dia in llamadas and (hoy + dt.timedelta(days=7)).isoformat() in llamadas and len(llamadas) == 2       # hoy y dentro de una semana (mismo día de la semana)
    # una cancha cambia sus horarios: se rehacen todas
    llamadas.clear()
    res = almacen.aplicar({"ops": [cancha(*c) for c in [("cn_f5a", "f5a", "Fútbol 5 · Cancha 1", "futbol", 60, 900, 1380), ("cn_f5b", "f5b", "Fútbol 5 · Cancha 2", "futbol", 60, 900, 1380),
                                                         ("cn_f7", "f7", "Fútbol 7 · Cancha completa", "futbol", 60, 900, 1380)]]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.escribir_ahora()
    assert sorted(llamadas) == sorted(ag.ventana(hoy))
    assert list(mapa(load_workbook(ad.ruta_agenda(dia))["Agenda"])["F"])[0] == "15 Hs"


def test_el_archivo_abierto_en_excel_no_rompe_nada_y_se_reintenta(almacen, rutas, monkeypatch):
    import predio.diario as diario

    hoy = dt.date.today()
    dia = hoy.isoformat()
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno", fecha=dia))
    reales = diario.os.replace

    def bloqueado(a, b):
        if str(b).endswith("_agenda.xlsx"):
            raise PermissionError("en uso")
        return reales(a, b)

    monkeypatch.setattr(diario.os, "replace", bloqueado)
    ad.marcar([dia])
    ad.escribir_ahora()                                                      # no levanta ningún error
    assert ad.ruta_excel(dia).exists() and ad.ruta_registro(dia).exists()    # el resumen y el registro salen igual
    assert not list(rutas.dias.rglob("*.tmp")) and not list(rutas.dias.glob("*/*_agenda.xlsx"))
    assert ad.agenda.fallidos() == ag.ventana(hoy) and ad.agenda.hay_pendiente() and "abierto" in ad.agenda.ultimo_error
    monkeypatch.setattr(diario.os, "replace", reales)
    ad.escribir_ahora()                                                      # se cerró el Excel: se escribe
    assert ad.agenda.fallidos() == [] and not ad.agenda.hay_pendiente() and ad.agenda.ultimo_error is None
    assert ad.ruta_agenda(dia).exists() and len(list(rutas.dias.glob("*/*_agenda.xlsx"))) == 8


def test_si_esta_abierta_la_de_hoy_las_demas_se_escriben(almacen, rutas, monkeypatch):
    import predio.diario as diario

    hoy = dt.date.today()
    ad = montar(almacen)
    reales = diario.os.replace
    monkeypatch.setattr(diario.os, "replace", lambda a, b: (_ for _ in ()).throw(PermissionError("en uso")) if str(b).endswith(f"{hoy.isoformat()}_agenda.xlsx") else reales(a, b))
    ad.marcar([hoy.isoformat()])
    ad.escribir_ahora()
    assert ad.agenda.fallidos() == [hoy.isoformat()] and len(list(rutas.dias.glob("*/*_agenda.xlsx"))) == 7


def esperar(condicion, segundos=10.0):
    limite = time.time() + segundos
    while time.time() < limite:
        if condicion():
            return True
        time.sleep(0.05)
    return False


def test_el_proceso_de_fondo_arma_las_agendas_solo_y_reintenta_si_estaba_abierta(almacen, rutas, monkeypatch):
    import predio.diario as diario

    hoy = dt.date.today()
    ad = montar(almacen)
    ad.ESPERA, ad.ESPERA_MAXIMA, ad.REINTENTO = 0.05, 0.2, 0.3
    reales = diario.os.replace
    abierto = {"si": True}

    def reemplazar(a, b):
        if abierto["si"] and str(b).endswith(f"{hoy.isoformat()}_agenda.xlsx"):
            raise PermissionError("en uso")
        return reales(a, b)

    monkeypatch.setattr(diario.os, "replace", reemplazar)
    ad.iniciar()                                                             # sin que nadie guarde nada: al arrancar se arman las agendas
    try:
        manana = (hoy + dt.timedelta(days=1)).isoformat()
        assert esperar(lambda: ad.ruta_agenda(manana).exists() and (hoy + dt.timedelta(days=7)).isoformat() in {p.parent.name for p in rutas.dias.glob("*/*_agenda.xlsx")})
        assert not ad.ruta_agenda(hoy.isoformat()).exists()
        abierto["si"] = False                                                # se cerró el Excel
        assert esperar(lambda: ad.ruta_agenda(hoy.isoformat()).exists(), 15)
        assert esperar(lambda: not ad.agenda.hay_pendiente())
    finally:
        ad.detener()


def test_al_empezar_otro_dia_se_arma_la_ventana_nueva(almacen, rutas):
    ad = montar(almacen)
    assert ad.agenda.cambio_de_dia(HOY) is True and ad.agenda.cambio_de_dia(HOY) is False       # la primera vez cuenta como "empezó el día"
    ad.agenda.escribir(HOY)
    assert ad.ruta_agenda("2026-10-22").exists() and not ad.ruta_agenda("2026-10-23").exists()
    assert ad.agenda.cambio_de_dia(HOY + dt.timedelta(days=1)) is True and ad.agenda.hay_pendiente()
    ad.agenda.escribir(HOY + dt.timedelta(days=1))
    assert ad.ruta_agenda("2026-10-23").exists() and ad.ruta_agenda(F).exists()                  # la de ayer queda como estaba
    assert not ad.agenda.hay_pendiente()


def test_lo_escrito_a_mano_en_plata_y_mp_se_conserva_al_regenerar(almacen):
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno"), turno("t2", "cn_pT", 1050, "Ana"))
    ws = hoja(ad)
    m = mapa(ws)
    wb = ws.parent
    ws.cell(m["F"]["19 Hs"], 2).value = 28000
    ws.cell(m["F"]["19 Hs"], 3).value = "QR"
    ws.cell(m["F"]["20 Hs"], 8).value = "=5000+2000"
    ws.cell(m["P"]["17:30 hs"], 9).value = "transf."
    ws.cell(m["F"]["14 Hs"], 4).value = "esto no se conserva (es una reserva)"
    wb.save(ad.ruta_agenda(F))
    res = almacen.aplicar({"ops": [turno("t1", "cn_f5a", 1140, "Bruno Cambiado")]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.agenda.escribir(HOY)
    ws = load_workbook(ad.ruta_agenda(F))["Agenda"]
    assert ws.cell(m["F"]["19 Hs"], 4).value == "Bruno Cambiado"                 # lo del programa se actualiza
    assert ws.cell(m["F"]["14 Hs"], 4).value is None                              # y lo que no era de las columnas $ y MP se rehace
    assert ws.cell(m["F"]["19 Hs"], 2).value == 28000 and ws.cell(m["F"]["19 Hs"], 3).value == "QR"
    assert ws.cell(m["F"]["20 Hs"], 8).value == "=5000+2000" and ws.cell(m["F"]["20 Hs"], 8).data_type == "f"      # una suma escrita a mano sigue siendo una suma
    assert ws.cell(m["P"]["17:30 hs"], 9).value == "transf."


def test_al_volver_a_abrir_el_programa_no_se_toca_lo_que_no_cambio(almacen, monkeypatch):
    """La firma de lo escrito queda en el archivo: un programa recién abierto no reescribe las agendas que siguen iguales (ni pisa lo que alguien anotó en ellas)."""
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno"))
    ws = hoja(ad)
    m = mapa(ws)
    ws.cell(m["F"]["14 Hs"], 4).value = "anotado a mano"          # no es una casilla de $ o MP: sólo se conserva si la agenda no cambia
    ws.cell(m["F"]["19 Hs"], 2).value = 28000
    ws.parent.save(ad.ruta_agenda(F))
    llamadas = contar_construcciones(monkeypatch)
    reabierto = ArchivosDiarios(almacen.db, almacen.rutas)                        # otro proceso: no sabe nada de lo escrito antes
    reabierto.agenda.pedir()
    reabierto.agenda.escribir(HOY)
    assert llamadas == []
    ws = load_workbook(ad.ruta_agenda(F))["Agenda"]
    assert ws.cell(m["F"]["14 Hs"], 4).value == "anotado a mano" and ws.cell(m["F"]["19 Hs"], 2).value == 28000
    # cambia un turno: se rehace esa agenda, conservando $ y MP
    res = almacen.aplicar({"ops": [turno("t1", "cn_f5a", 1140, "Bruno 2")]}, "Caja", None)
    reabierto.marcar(res["dias"])
    reabierto.agenda.escribir(HOY)
    assert llamadas == [F]
    ws = load_workbook(ad.ruta_agenda(F))["Agenda"]
    assert ws.cell(m["F"]["19 Hs"], 4).value == "Bruno 2" and ws.cell(m["F"]["19 Hs"], 2).value == 28000 and ws.cell(m["F"]["14 Hs"], 4).value is None


def test_si_el_archivo_esta_bloqueado_para_leerlo_no_se_pisa_y_se_reintenta(almacen, monkeypatch):
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno"))
    ws = hoja(ad)
    ws.cell(mapa(ws)["F"]["19 Hs"], 2).value = 28000
    ws.parent.save(ad.ruta_agenda(F))
    real = ag.load_workbook
    bloqueado = {"si": True}

    def abrir(ruta, *a, **k):
        if bloqueado["si"] and str(ruta).endswith(f"{F}_agenda.xlsx"):
            raise PermissionError("en uso")
        return real(ruta, *a, **k)

    monkeypatch.setattr(ag, "load_workbook", abrir)
    res = almacen.aplicar({"ops": [turno("t1", "cn_f5a", 1140, "Bruno 2")]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.agenda.escribir(HOY)
    assert ad.agenda.fallidos() == [F] and "abierto" in ad.agenda.ultimo_error
    assert load_workbook(ad.ruta_agenda(F))["Agenda"].cell(mapa(ws)["F"]["19 Hs"], 2).value == 28000        # el archivo quedó como estaba: no se perdió lo anotado
    bloqueado["si"] = False
    ad.agenda.escribir(HOY)
    ws = load_workbook(ad.ruta_agenda(F))["Agenda"]
    assert ad.agenda.fallidos() == [] and ws.cell(mapa(ws)["F"]["19 Hs"], 4).value == "Bruno 2" and ws.cell(mapa(ws)["F"]["19 Hs"], 2).value == 28000


def test_si_no_se_pueden_leer_los_turnos_se_reintenta_sin_darse_vueltas(almacen, monkeypatch):
    ad = montar(almacen)
    real = ag.leer_datos

    def falla(*a, **k):
        raise sqlite3.OperationalError("database is locked")

    monkeypatch.setattr(ag, "leer_datos", falla)
    ad.agenda.pedir()
    ad.agenda.escribir(HOY)
    assert ad.agenda.fallidos() == ag.ventana(HOY) and not ad.agenda.por_revisar() and "locked" in ad.agenda.ultimo_error      # queda para reintentar cada tanto, no en el acto
    monkeypatch.setattr(ag, "leer_datos", real)
    ad.agenda.escribir(HOY)
    assert ad.agenda.fallidos() == [] and ad.ruta_agenda(F).exists() and ad.agenda.ultimo_error is None


def test_las_canchas_de_papel_son_las_mismas_en_la_pantalla_la_nube_y_la_agenda():
    """Los horarios de papel están escritos en la pantalla (CANCHAS_PAPEL), en la nube (schema.sql) y como respaldo en la agenda: no pueden diferir."""
    html = (RAIZ / "predio" / "web" / "index.html").read_text(encoding="utf-8")
    objeto = re.search(r"const CANCHAS_PAPEL = (\{.*?\n\});", html, re.S).group(1)
    r = subprocess.run(["node", "-e", f"console.log(JSON.stringify(({objeto})))"], capture_output=True, text=True, timeout=30) if shutil.which("node") else None
    if r is None:
        pytest.skip("falta node")
    pantalla = json.loads(r.stdout)
    nube = sqlite3.connect(":memory:")
    nube.executescript((RAIZ / "nube" / "schema.sql").read_text(encoding="utf-8"))
    de_la_nube = {c: {"nombre": n, "ap": ap, "ci": ci, "dur": dur} for c, n, dur, ap, ci in nube.execute("SELECT id, nombre, dur, ap, ci FROM canchas")}
    assert pantalla == de_la_nube
    for clave, c in pantalla.items():
        assert (c["ap"], c["ci"], c["dur"]) == ag.HORARIOS_PAPEL["futbol" if clave.startswith("f") else "padel"]


def test_un_archivo_viejo_danado_no_frena_la_agenda_nueva(almacen):
    ad = montar(almacen, turno("t1", "cn_f5a", 1140, "Bruno"))
    hoja(ad)
    ad.ruta_agenda(F).write_bytes(b"esto no es un excel")
    res = almacen.aplicar({"ops": [turno("t1", "cn_f5a", 1140, "Nuevo")]}, "Caja", None)
    ad.marcar(res["dias"])
    ad.agenda.escribir(HOY)
    ws = load_workbook(ad.ruta_agenda(F))["Agenda"]
    assert reservas(ws, mapa(ws)["F"]["19 Hs"])[0] == "Nuevo"


def test_si_borran_el_archivo_se_vuelve_a_escribir(almacen):
    ad = montar(almacen)
    hoja(ad)
    ad.ruta_agenda(F).unlink()
    ad.agenda.pedir()
    ad.agenda.escribir(HOY)
    assert ad.ruta_agenda(F).exists()


# ------------------------------------------------------------------ equivalencia con la pantalla

def _js_ocurrencias(canchas, fijos, turnos, fechas):
    html = (RAIZ / "predio" / "web" / "index.html").read_text(encoding="utf-8")
    cancha_js = re.search(r"^const cancha = .*;$", html, re.M).group(0)
    dow_js = re.search(r"^const dowDe = .*;$", html, re.M).group(0)
    ini = html.index("function ocurrenciasFijas(fecha) {")
    funcion = html[ini:html.index("\n}\n", ini) + 2]
    script = (
        "const vm = require('vm'); const e = JSON.parse(require('fs').readFileSync(0, 'utf8'));\n"
        "const ctx = { DB: { canchas: e.canchas, fijos: e.fijos, turnos: e.turnos }, fechas: e.fechas, out: null }; vm.createContext(ctx);\n"
        "vm.runInContext(e.codigo + '\\nout = fechas.map(f => ocurrenciasFijas(f).map(t => [t.fijoId, t.canchaId, t.ini, t.fin, t.cliente, t.nota]));', ctx);\n"
        "console.log(JSON.stringify(ctx.out));\n")
    r = subprocess.run(["node", "-e", script], input=json.dumps({"codigo": "\n".join([cancha_js, dow_js, funcion]), "canchas": canchas, "fijos": fijos, "turnos": turnos, "fechas": fechas}),
                       capture_output=True, text=True, timeout=60)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


@pytest.mark.skipif(not shutil.which("node"), reason="falta node")
def test_las_ocurrencias_de_fijos_son_las_mismas_que_calcula_la_pantalla():
    canchas = [{"id": "a", "dur": 60}, {"id": "b", "dur": 90}]
    fijos = [{"id": "f1", "canchaId": "a", "dow": 4, "ini": 840, "cliente": "Uno", "nota": "n", "desde": "2026-09-01", "hasta": None},
             {"id": "f2", "canchaId": "b", "dow": 4, "ini": 600, "cliente": "Dos", "nota": "", "desde": "2026-10-22", "hasta": None},
             {"id": "f3", "canchaId": "a", "dow": 4, "ini": 900, "cliente": "Tres", "nota": "", "desde": "2026-01-01", "hasta": "2026-10-14"},
             {"id": "f4", "canchaId": "a", "dow": 4, "ini": 960, "cliente": "Cuatro", "nota": "", "desde": "2026-01-01", "hasta": "2026-10-15"},
             {"id": "f5", "canchaId": "b", "dow": 0, "ini": 1320, "cliente": "Cinco", "nota": "", "desde": "2026-01-01", "hasta": ""},
             {"id": "f6", "canchaId": "a", "dow": 1, "ini": 1020, "cliente": "Seis", "nota": "x", "desde": "2026-10-12", "hasta": "2026-10-26"},
             {"id": "f7", "canchaId": "b", "dow": 6, "ini": 690, "cliente": "Siete", "nota": "", "desde": "2026-10-17", "hasta": None}]
    turnos = [{"id": "t1", "fijoId": "f1", "ocurrencia": "2026-10-15", "fecha": "2026-10-15", "estado": "cancelado"},      # suspendido
              {"id": "t2", "fijoId": "f4", "ocurrencia": "2026-10-15", "fecha": "2026-10-16", "estado": "reservado"},     # movido
              {"id": "t3", "fijoId": "f1", "ocurrencia": "2026-10-22", "fecha": "2026-10-22", "estado": "cobrado"},       # ya cobrado
              {"id": "t4", "fijoId": None, "ocurrencia": None, "fecha": "2026-10-15", "estado": "reservado"},
              {"id": "t5", "fijoId": "f5", "ocurrencia": "2026-10-18", "fecha": "2027-01-01", "estado": "reservado"},
              {"id": "t6", "fijoId": "f6", "ocurrencia": "2026-10-19", "fecha": "2026-10-19", "estado": "reservado"}]
    fechas = [(dt.date(2026, 10, 8) + dt.timedelta(days=i)).isoformat() for i in range(0, 40)]
    de_js = _js_ocurrencias(canchas, fijos, turnos, fechas)
    por_id = {c["id"]: c for c in canchas}
    de_py = [[[t["fijoId"], t["canchaId"], t["ini"], t["fin"], t["cliente"], t["nota"]] for t in ag.ocurrencias_fijas(f, fijos, turnos, por_id)] for f in fechas]
    assert de_py == de_js
    assert sum(len(x) for x in de_py) > 15 and any(len(x) == 0 for x in de_py)       # la prueba no es trivial


# ------------------------------------------------------------------ horarios reales en la nube y con la API

def test_las_semillas_de_la_nube_y_su_migracion_dejan_los_mismos_horarios(tmp_path):
    sql_schema = (RAIZ / "nube" / "schema.sql").read_text(encoding="utf-8")
    sql_migracion = (RAIZ / "nube" / "migraciones.sql").read_text(encoding="utf-8")
    esperado = {"f5a": ("Fútbol 5 · Cancha 1", 60, 840, 1380), "f5b": ("Fútbol 5 · Cancha 2", 60, 840, 1380), "f7": ("Fútbol 7 · Cancha completa", 60, 840, 1380),
                "padelT": ("Pádel · Cancha 1", 90, 600, 1410), "padelA": ("Pádel · Cancha 2", 90, 600, 1410)}
    nuevo = sqlite3.connect(":memory:")
    nuevo.executescript(sql_schema)
    assert {r[0]: tuple(r[1:]) for r in nuevo.execute("SELECT id, nombre, dur, ap, ci FROM canchas")} == esperado
    # una base de la nube que ya tenía las canchas de ejemplo (con sus precios): la migración cambia horarios y nombres y no toca los precios
    viejo = sqlite3.connect(":memory:")
    viejo.executescript(sql_schema.split("-- Horarios y nombres de la planilla de papel")[0])
    viejo.executescript("""INSERT INTO canchas VALUES ('f5a','f5a','Fútbol 5 · Cancha 1','futbol',60,600,1440,31000,26000,1), ('f5b','f5b','Fútbol 5 · Cancha 2','futbol',60,600,1440,31000,26000,2),
        ('f7','f7','Fútbol 7 · Cancha completa','futbol',60,600,1440,50000,45000,3), ('padelT','padelT','Pádel · Techada','padel',90,480,1440,27000,23000,4),
        ('padelA','padelA','Pádel · Abierta','padel',90,480,1440,23000,20000,5);""")
    viejo.executescript(sql_migracion)
    assert {r[0]: tuple(r[1:]) for r in viejo.execute("SELECT id, nombre, dur, ap, ci FROM canchas")} == esperado
    assert {r[0]: (r[1], r[2]) for r in viejo.execute("SELECT id, precio, precio_fijo FROM canchas")}["padelT"] == (27000, 23000)
    assert {r[0]: (r[1], r[2]) for r in viejo.execute("SELECT id, precio, precio_fijo FROM canchas")}["f5a"] == (31000, 26000)


def test_la_agenda_sale_sola_despues_de_guardar_por_la_api(srv):
    hoy = dt.date.today()
    base = base_inicial()
    base["ops"] += [cancha(*c) for c in CANCHAS] + [turno("t1", "cn_f5a", 1140, "Desde la pantalla", fecha=hoy.isoformat())]
    assert pedir(srv, "POST", "/api/guardar", {**base, "base_rev": 0, "usuario": "Marcos"})[0] == 200
    diario = srv.app.diario
    assert esperar(lambda: diario.ruta_agenda(hoy.isoformat()).exists() and diario.ruta_agenda((hoy + dt.timedelta(days=7)).isoformat()).exists())
    ws = load_workbook(diario.ruta_agenda(hoy.isoformat()))["Agenda"]
    assert reservas(ws, mapa(ws)["F"]["19 Hs"]) == ["Desde la pantalla", "5", None, None]
    salud = pedir(srv, "GET", "/api/salud")[1]
    assert salud["diario"]["agenda"]["pendientes"] == [] and salud["diario"]["agenda"]["error"] is None
