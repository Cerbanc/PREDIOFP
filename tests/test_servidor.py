import io
import json
import threading
import urllib.error
import urllib.request

import pytest
from openpyxl import Workbook, load_workbook

from conftest import base_inicial, op
from predio import app as app_mod
from predio.config import AJUSTES_POR_DEFECTO
from predio.servidor import Servidor


@pytest.fixture()
def srv(rutas, monkeypatch):
    abiertas = []
    monkeypatch.setattr("predio.servidor.abrir_carpeta", lambda ruta: abiertas.append(str(ruta)))
    app = app_mod.crear_aplicacion(rutas, dict(AJUSTES_POR_DEFECTO, copias_conservar_dias=30), "real")
    app.diario.ESPERA, app.diario.ESPERA_MAXIMA = 0.05, 0.2
    app.diario.iniciar()
    s = Servidor(app, 0)
    s.iniciar()
    s.abiertas = abiertas
    yield s
    s.parar()
    app_mod.cerrar_aplicacion(app)


def pedir(s, metodo, ruta, cuerpo=None, token=True, cabeceras=None, crudo=False, host=None):
    url = f"http://127.0.0.1:{s.puerto}{ruta}"
    datos = None if cuerpo is None else (cuerpo if isinstance(cuerpo, bytes) else json.dumps(cuerpo).encode())
    h = {"Host": host or f"127.0.0.1:{s.puerto}"}
    if token:
        h["X-Predio-Token"] = s.app.token
    if datos is not None and not isinstance(cuerpo, bytes):
        h["Content-Type"] = "application/json"
    h.update(cabeceras or {})
    req = urllib.request.Request(url, data=datos, method=metodo, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            b = r.read()
            return r.status, (b if crudo else json.loads(b or b"null")), r.headers
    except urllib.error.HTTPError as e:
        b = e.read()
        try:
            return e.code, json.loads(b), e.headers
        except ValueError:
            return e.code, b, e.headers


def test_la_pagina_lleva_la_clave_y_cabeceras_de_seguridad(srv):
    estado, html, cab = pedir(srv, "GET", "/", token=False, crudo=True)
    assert estado == 200 and b"window.PREDIO=" in html and srv.app.token.encode() in html
    assert b"fonts.googleapis" not in html and b"cdnjs" not in html
    assert "default-src 'self'" in cab["Content-Security-Policy"] and cab["X-Frame-Options"] == "DENY"


def test_ping_no_pide_clave_pero_lo_demas_si(srv):
    assert pedir(srv, "GET", "/api/ping", token=False)[1]["app"] == "predio"
    estado, cuerpo, _ = pedir(srv, "GET", "/api/estado", token=False)
    assert estado == 403 and "recargala" in cuerpo["error"]
    estado, _, _ = pedir(srv, "POST", "/api/guardar", {"ops": []}, token=False)
    assert estado == 403


def test_otra_web_no_puede_tocar_los_datos(srv):
    estado, cuerpo, _ = pedir(srv, "GET", "/api/estado", host="evil.example.com")
    assert estado == 403
    estado, cuerpo, _ = pedir(srv, "POST", "/api/guardar", {"ops": []}, cabeceras={"Origin": "https://evil.example.com"})
    assert estado == 403


def test_estado_vacio_y_guardado_completo(srv):
    e = pedir(srv, "GET", "/api/estado")[1]
    assert e["vacia"] is True and e["modo"] == "real" and e["rev"] == 0
    estado, r, _ = pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0, "usuario": "Sistema"})
    assert estado == 200 and r["rev"] == 1
    e = pedir(srv, "GET", "/api/estado")[1]
    assert e["vacia"] is False and len(e["db"]["productos"]) == 2 and e["rev"] == 1


def test_ventana_vieja_recibe_409(srv):
    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0})
    pedir(srv, "POST", "/api/guardar", {"ops": [op("negocio", "_", "Uno")], "base_rev": 1})
    estado, r, _ = pedir(srv, "POST", "/api/guardar", {"ops": [op("negocio", "_", "Dos")], "base_rev": 1})
    assert estado == 409 and r["conflicto"] is True and r["rev"] == 2


def test_guardado_invalido_es_400_y_no_rompe(srv):
    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0})
    estado, r, _ = pedir(srv, "POST", "/api/guardar", {"ops": [{"c": "nada", "id": "x", "d": {}}], "base_rev": 1})
    assert estado == 400 and "desconocida" in r["error"]
    assert pedir(srv, "GET", "/api/estado")[1]["rev"] == 1


def test_el_excel_del_dia_se_genera_solo_despues_de_guardar(srv):
    import time

    from predio.util import ahora_ms, dia_de

    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0})
    dia = dia_de(ahora_ms())
    for _ in range(80):
        if srv.app.diario.ruta_excel(dia).exists():
            break
        time.sleep(0.1)
    assert srv.app.diario.ruta_excel(dia).exists() and srv.app.diario.ruta_registro(dia).exists()


def test_auditoria_y_dia(srv):
    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0, "usuario": "Marcos"})
    r = pedir(srv, "GET", "/api/auditoria?importantes=1")[1]
    assert r["ok"] and r["filas"] and all(f["importante"] == 1 for f in r["filas"])
    det = pedir(srv, "GET", f"/api/auditoria/{r['filas'][0]['id']}")[1]["fila"]
    assert det["usuario"] == "Marcos"
    assert pedir(srv, "GET", "/api/auditoria/99999")[0] == 404
    d = pedir(srv, "GET", f"/api/dia/{__import__('predio.util', fromlist=['x']).dia_de(__import__('predio.util', fromlist=['x']).ahora_ms())}")[1]
    assert d["ok"] and "kpis" in d and "ventas" not in d
    assert pedir(srv, "GET", "/api/dia/ayer")[0] == 404


def test_copias_por_la_api(srv):
    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0})
    r = pedir(srv, "POST", "/api/copias/crear", {})[1]
    assert r["ok"] and r["archivo"].startswith("predio_")
    assert any(c["archivo"] == r["archivo"] for c in pedir(srv, "GET", "/api/copias")[1]["copias"])
    pedir(srv, "POST", "/api/guardar", {"ops": [op("negocio", "_", "Cambiado")], "base_rev": 1})
    res = pedir(srv, "POST", "/api/copias/restaurar", {"archivo": r["archivo"], "usuario": "Marcos"})
    assert res[0] == 200
    e = pedir(srv, "GET", "/api/estado")[1]
    assert e["db"]["negocio"] == "Predio Deportivo" and e["rev"] > 2
    assert pedir(srv, "POST", "/api/copias/restaurar", {"archivo": "../x.db"})[0] == 400


def test_salud_y_abrir_carpeta(srv):
    s = pedir(srv, "GET", "/api/salud")[1]
    assert s["ok"] and s["carpeta"] == str(srv.app.rutas.raiz) and s["modo"] == "real"
    assert pedir(srv, "POST", "/api/abrir-carpeta", {"que": "copias"})[0] == 200
    assert pedir(srv, "POST", "/api/abrir-carpeta", {"que": "dia"})[0] == 200
    assert pedir(srv, "POST", "/api/abrir-carpeta", {"que": "../../etc"})[0] == 400
    assert str(srv.app.rutas.copias) in srv.abiertas and any("dias" in a for a in srv.abiertas)


def test_planilla_plantilla_exportar_y_leer(srv):
    pedir(srv, "POST", "/api/guardar", {**base_inicial(), "base_rev": 0})
    estado, datos, cab = pedir(srv, "GET", "/api/planilla/exportar.xlsx", crudo=True)
    assert estado == 200 and "attachment" in cab["Content-Disposition"]
    wb = load_workbook(io.BytesIO(datos))
    nombres = [fila[1].value for fila in wb["Productos"].iter_rows(min_row=2) if fila[1].value]
    assert nombres == ["Pancho", "Coca-Cola 500 ml"]      # ordenado por categoría (Buffet, Kiosco) y nombre
    assert any(p.name.startswith("productos_y_stock_") for p in srv.app.rutas.planillas.iterdir())
    estado, datos, _ = pedir(srv, "GET", "/api/planilla/plantilla.xlsx", crudo=True)
    assert estado == 200 and load_workbook(io.BytesIO(datos))["Productos"]["B2"].value is None
    # leer una planilla completada
    wb = Workbook(); ws = wb.active; ws.append(["Producto", "Categoría", "Precio", "Stock"]); ws.append(["Alfajor", "Kiosco", 1200, 40])
    buf = io.BytesIO(); wb.save(buf)
    estado, r, _ = pedir(srv, "POST", "/api/planilla/leer", buf.getvalue(), cabeceras={"X-Nombre": "lista%20nueva.xlsx", "Content-Type": "application/octet-stream"})
    assert estado == 200 and r["filas"][0]["nombre"] == "Alfajor" and r["filas"][0]["stock"] == 40
    estado, r, _ = pedir(srv, "POST", "/api/planilla/leer", b"basura", cabeceras={"X-Nombre": "x.xlsx"})
    assert estado == 400 and "Excel" in r["error"]


def test_imagenes_servidas_desde_archivo(srv):
    import base64

    png = "data:image/png;base64," + base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"1" * 40).decode()
    base = base_inicial()
    for o in base["ops"]:
        if o["id"] == "p1":
            o["d"]["img"] = png
    r = pedir(srv, "POST", "/api/guardar", {**base, "base_rev": 0})[1]
    url = r["reemplazos"][0]["valor"]
    estado, datos, cab = pedir(srv, "GET", url, token=False, crudo=True)
    assert estado == 200 and cab["Content-Type"] == "image/png" and datos.startswith(b"\x89PNG")
    assert pedir(srv, "GET", "/img/../../datos/predio.db", token=False)[0] == 404
    assert pedir(srv, "GET", "/img/aaaaaaaaaaaaaaaaaaaa.png", token=False)[0] == 404


def test_cerrar_pide_apagar(srv):
    assert pedir(srv, "POST", "/api/cerrar", {})[1]["ok"]
    assert srv.app.apagar.wait(3)


def test_imagen_desde_link_rechaza_cosas_raras(srv):
    assert pedir(srv, "POST", "/api/imagen-desde-link", {"url": "file:///etc/passwd"})[0] == 400
    assert pedir(srv, "POST", "/api/imagen-desde-link", {"url": "ftp://x/y.png"})[0] == 400
