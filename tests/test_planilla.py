import io

import pytest
from openpyxl import Workbook, load_workbook

from predio.planilla import COLUMNAS, PlanillaInvalida, construir_planilla, leer_planilla, numero_ar

PRODUCTOS = [
    {"id": "p1", "nombre": "Coca-Cola 500 ml", "catId": "c1", "tipo": "directo", "precio": 2200, "costo": 1400, "stock": 8, "minimo": 6, "paqueteNombre": "Bulto 6 x 12", "paqueteUnidades": 72, "activo": True, "receta": [], "link": ""},
    {"id": "p2", "nombre": "Salchicha", "catId": "c2", "tipo": "insumo", "precio": 0, "costo": 350, "stock": 24, "minimo": 6, "paqueteUnidades": 6, "activo": True, "receta": []},
    {"id": "p3", "nombre": "Pan de pancho", "catId": "c2", "tipo": "insumo", "precio": 0, "costo": 250, "stock": 30, "minimo": 6, "paqueteUnidades": 1, "activo": True, "receta": []},
    {"id": "p4", "nombre": "Pancho", "catId": "c2", "tipo": "receta", "precio": 3500, "costo": 0, "stock": 0, "minimo": 0, "activo": False, "receta": [{"id": "p2", "cant": 1}, {"id": "p3", "cant": 1}]},
    {"id": "p5", "nombre": "Café", "catId": "c2", "tipo": "servicio", "precio": 1800, "costo": 500, "stock": 0, "minimo": 0, "activo": True, "receta": []},
]
CATEGORIAS = [{"id": "c1", "nombre": "Kiosco"}, {"id": "c2", "nombre": "Buffet"}]


def test_exportar_y_leer_de_vuelta():
    data = construir_planilla(PRODUCTOS, CATEGORIAS)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Instrucciones", "Productos", "Listas"]
    r = leer_planilla(data, "productos.xlsx")
    assert r["columnas"][0] == "nombre" and set(r["columnas"]) == {c[0] for c in COLUMNAS}
    por_nombre = {f["nombre"]: f for f in r["filas"]}
    assert len(por_nombre) == 5
    coca = por_nombre["Coca-Cola 500 ml"]
    assert (coca["id"], coca["categoria"], coca["tipo"], coca["precio"], coca["costo"], coca["stock"], coca["minimo"], coca["unidades"], coca["paquete"], coca["activo"]) == \
        ("p1", "Kiosco", "directo", 2200, 1400, 8, 6, 72, "Bulto 6 x 12", True)
    pancho = por_nombre["Pancho"]
    assert pancho["tipo"] == "receta" and pancho["receta"] == "1 Salchicha + 1 Pan de pancho" and pancho["activo"] is False
    assert por_nombre["Salchicha"]["precio"] is None and por_nombre["Café"]["tipo"] == "servicio"
    assert por_nombre["Café"]["stock"] is None


def test_la_plantilla_vacia_tiene_los_titulos_y_las_listas():
    data = construir_planilla([], CATEGORIAS, con_datos=False)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Productos"]
    assert [c.value for c in ws[1]] == [c[1] for c in COLUMNAS]
    assert ws.max_row >= 400 and ws["B2"].value is None
    assert [r[0].value for r in wb["Listas"].iter_rows(min_row=2, max_row=3)] == ["Kiosco", "Buffet"]
    with pytest.raises(PlanillaInvalida):
        leer_planilla(data, "plantilla.xlsx")      # sin productos todavía


def test_leer_un_excel_cualquiera_con_titulos_distintos():
    wb = Workbook()
    ws = wb.active
    ws.append(["Lista de precios del kiosco"])
    ws.append([])
    ws.append(["Artículo", "Rubro", "Precio", "Cantidad", "Mínimo"])
    ws.append(["Alfajor", "Kiosco", "$ 1.200", 40, 10])
    ws.append(["Papas fritas", "Kiosco", "2000,50", "15", None])
    ws.append([None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = leer_planilla(buf.getvalue(), "lista.xlsx")
    assert [f["nombre"] for f in r["filas"]] == ["Alfajor", "Papas fritas"]
    assert r["filas"][0]["precio"] == 1200 and r["filas"][0]["stock"] == 40 and r["filas"][0]["minimo"] == 10
    assert r["filas"][1]["precio"] == 2001 and r["filas"][1]["stock"] == 15 and r["filas"][1]["minimo"] is None
    assert "id" not in r["columnas"] and "costo" not in r["columnas"]
    assert r["filas"][0]["n"] == 4


def test_csv_con_punto_y_coma_y_acentos_de_windows():
    texto = "Producto;Categoría;Precio;Stock\nCerveza 1 L;Buffet;4.200;18\nPorción de muzzarella;Buffet;2800;16\n"
    r = leer_planilla(texto.encode("cp1252"), "export.csv")
    assert [f["nombre"] for f in r["filas"]] == ["Cerveza 1 L", "Porción de muzzarella"]
    assert r["filas"][0]["precio"] == 4200 and r["filas"][0]["categoria"] == "Buffet"


def test_archivos_invalidos():
    with pytest.raises(PlanillaInvalida):
        leer_planilla(b"esto no es excel", "x.xlsx")
    with pytest.raises(PlanillaInvalida):
        leer_planilla(b"abc", "x.pdf")
    with pytest.raises(PlanillaInvalida):
        leer_planilla("Precio;Stock\n1;2\n".encode(), "x.csv")


@pytest.mark.parametrize("entrada,esperado", [("2200", 2200), ("$ 2.200", 2200), ("2.200,50", 2201), ("1,5", 2), (1500.4, 1500), (None, None), ("", None), ("abc", None), ("-5", -5), (True, 1)])
def test_numero_ar(entrada, esperado):
    assert numero_ar(entrada) == esperado


def test_stock_minimo_no_se_confunde_con_stock_en_ningun_orden():
    for cab in (["Producto", "Stock actual", "Stock mínimo"], ["Producto", "Stock mínimo", "Stock actual"], ["Producto", "Stock", "Stock mínimo"]):
        texto = ";".join(cab) + "\nCoca;36;6\n" if cab[1] != "Stock mínimo" else ";".join(cab) + "\nCoca;6;36\n"
        f = leer_planilla(texto.encode(), "x.csv")["filas"][0]
        assert (f["stock"], f["minimo"]) == (36, 6), cab
