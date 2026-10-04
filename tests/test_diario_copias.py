import datetime as dt
import sqlite3
import time

import pytest
from openpyxl import load_workbook

from conftest import base_inicial, borrar, op
from predio.copias import Copias, CopiaInvalida
from predio.diario import ArchivosDiarios, ArchivoEnUso, datos_del_dia, texto_registro
from predio.util import ahora_ms, dia_de


def venta(id_, nro, ts, total, destino="Mesa 1", pagos=None, estado="ok", items=None, descuento=0, ajuste=None):
    return {"id": id_, "nro": nro, "ts": ts, "destino": destino, "subtotal": total + descuento, "descuento": descuento, "ajuste": ajuste, "total": total, "estado": estado,
            "items": items or [{"prodId": "p1", "nombre": "Coca-Cola 500 ml", "catId": "c1", "cant": 2, "precio": total // 2, "costo": 1400, "nota": ""}],
            "pagos": pagos or [{"metodoId": "mp1", "nombre": "Efectivo", "monto": total}]}


@pytest.fixture()
def con_datos(almacen):
    almacen.aplicar(base_inicial(), "Caja", 0)
    ts = ahora_ms()
    almacen.aplicar({"ops": [
        op("cajas", "cx1", {"id": "cx1", "abierta": ts - 3_600_000, "fondo": 15000, "responsable": "Marcos", "cerrada": None}),
        op("ventas", "v1", venta("v1", 1, ts - 600_000, 4400)),
        op("ventas", "v2", venta("v2", 2, ts - 300_000, 3500, destino="Barra", pagos=[{"metodoId": "mp2", "nombre": "MP · QR", "monto": 3500}], descuento=500, ajuste={"tipo": "monto", "motivo": "Promo", "nota": ""})),
        op("ventas", "v3", venta("v3", 3, ts - 200_000, 1000, estado="anulada")),
        op("caja", "cj1", {"id": "cj1", "ts": ts - 600_000, "metodoId": "mp1", "nombre": "Efectivo", "monto": 4400, "tipo": "cobro", "motivo": "Ticket #1"}),
        op("caja", "cj2", {"id": "cj2", "ts": ts - 300_000, "metodoId": "mp2", "nombre": "MP · QR", "monto": 3500, "tipo": "cobro", "motivo": "Ticket #2"}),
        op("caja", "cj3", {"id": "cj3", "ts": ts - 100_000, "metodoId": "mp1", "nombre": "Efectivo", "monto": -500, "tipo": "devolucion", "motivo": "Devolución"}),
        op("movCaja", "mc1", {"id": "mc1", "ts": ts - 50_000, "cajaId": "cx1", "tipo": "gasto", "concepto": "Limpieza", "proveedor": "", "monto": 2500, "nota": ""}),
        op("movimientos", "mv1", {"id": "mv1", "ts": ts - 600_000, "prodId": "p1", "nombre": "Coca-Cola 500 ml", "tipo": "venta", "cant": -2, "antes": 12, "despues": 10, "motivo": "Ticket #1"}),
    ]}, "Marcos")
    return almacen


def test_datos_del_dia_calcula_totales(con_datos):
    con = con_datos.db.lectura()
    d = datos_del_dia(con, dia_de(ahora_ms()))
    con.close()
    k = d["kpis"]
    assert (k["tickets"], k["anulados"], k["total_vendido"], k["descuentos"]) == (2, 1, 7900, 500)
    medios = {m["medio"]: m for m in d["por_medio"]}
    assert medios["Efectivo"]["neto"] == 3900 and medios["MP · QR"]["neto"] == 3500
    assert k["cobrado_neto"] == 7400
    assert d["movimientos_caja"]["gasto"] == 2500
    assert d["por_categoria"][0]["categoria"] == "Kiosco"
    assert {x["destino"] for x in d["por_destino"]} == {"Mesa 1", "Barra"}
    assert d["kpis"]["cambios_importantes"] >= 3


def test_excel_y_registro_del_dia(con_datos, rutas):
    dia = dia_de(ahora_ms())
    ad = ArchivosDiarios(con_datos.db, rutas)
    ad.generar_dia(dia)
    wb = load_workbook(ad.ruta_excel(dia))
    assert wb.sheetnames == ["Resumen", "Ventas", "Detalle", "Libro de caja", "Movimientos de caja", "Stock", "Turnos", "Cambios"]
    textos = [c.value for fila in wb["Resumen"].iter_rows() for c in fila if c.value is not None]
    assert "Total vendido" in textos and 7900 in textos
    ventas = list(wb["Ventas"].iter_rows(min_row=2, values_only=True))
    assert [v[0] for v in ventas] == [1, 2, 3]
    assert ventas[1][9] == "Promo"
    assert len(list(wb["Detalle"].iter_rows(min_row=2))) == 3
    log = ad.ruta_registro(dia).read_text(encoding="utf-8")
    assert "Ticket #2 cobrado" in log and "CON DESCUENTO" in log and "!" in log
    assert not list(ad.carpeta(dia).glob("*.tmp"))


def test_el_excel_abierto_no_rompe_y_se_reintenta(con_datos, rutas, monkeypatch):
    import predio.diario as diario

    dia = dia_de(ahora_ms())
    ad = ArchivosDiarios(con_datos.db, rutas)
    reales = diario.os.replace

    def bloqueado(a, b):
        if str(b).endswith(".xlsx"):
            raise PermissionError("en uso")
        return reales(a, b)

    monkeypatch.setattr(diario.os, "replace", bloqueado)
    with pytest.raises(ArchivoEnUso):
        ad.generar_dia(dia)
    assert not list(ad.carpeta(dia).glob("*.tmp"))
    monkeypatch.setattr(diario.os, "replace", reales)
    ad.generar_dia(dia)
    assert ad.ruta_excel(dia).exists()


def test_el_proceso_de_fondo_escribe_solo(con_datos, rutas):
    dia = dia_de(ahora_ms())
    ad = ArchivosDiarios(con_datos.db, rutas)
    ad.ESPERA, ad.ESPERA_MAXIMA = 0.1, 0.5
    ad.iniciar()
    try:
        ad.marcar([dia])
        for _ in range(60):
            if ad.ruta_excel(dia).exists():
                break
            time.sleep(0.1)
        assert ad.ruta_excel(dia).exists() and ad.ruta_registro(dia).exists()
        assert ad.pendientes() == []
    finally:
        ad.detener()


def test_copia_y_restauracion(con_datos, rutas):
    copias = Copias(con_datos.db, con_datos, rutas)
    copia = copias.crear("manual")
    assert copia.exists() and copias.listar()[0]["motivo"] == "manual"
    # cambios posteriores a la copia
    con_datos.aplicar({"ops": [op("negocio", "_", "Cambiado")]}, "Marcos")
    con_datos.aplicar({"ops": [borrar("productos", "p1")]}, "Marcos")
    rev_antes = con_datos.rev()
    r = copias.restaurar(copia.name, "Marcos")
    assert r["rev"] > rev_antes
    import json

    e = json.loads(con_datos.estado_json())
    assert e["db"]["negocio"] == "Predio Deportivo"
    assert [p["id"] for p in e["db"]["productos"]] == ["p1", "p2"]
    # quedó una copia de seguridad previa y el aviso en el registro
    assert any(c["motivo"].startswith("seguridad") for c in copias.listar())
    ult = con_datos.db.con.execute("SELECT resumen, importante FROM auditoria ORDER BY id DESC LIMIT 1").fetchone()
    assert "Se restauró la copia" in ult[0] and ult[1] == 1
    # la base sigue funcionando después de restaurar
    con_datos.aplicar({"ops": [op("negocio", "_", "Seguimos")]}, "Marcos", r["rev"])


def test_restaurar_rechaza_nombres_raros(con_datos, rutas):
    copias = Copias(con_datos.db, con_datos, rutas)
    for malo in ("../datos/predio.db", "predio_x.db", "", "predio_2026-01-01_0000_a/../../x.db"):
        with pytest.raises(CopiaInvalida):
            copias.restaurar(malo)


def test_limpieza_de_copias_viejas(con_datos, rutas):
    copias = Copias(con_datos.db, con_datos, rutas, conservar_dias=30)
    hoy = dt.date.today()
    nombres = []
    for dias_atras, hhmm in [(0, "0900"), (0, "1800"), (10, "0900"), (10, "2100"), (40, "0900"), (100, "0900")]:
        n = f"predio_{(hoy - dt.timedelta(days=dias_atras)).isoformat()}_{hhmm}_inicio.db"
        (rutas.copias / n).write_bytes(b"x")
        nombres.append(n)
    copias.limpiar()
    quedan = {c["archivo"] for c in copias.listar()}
    assert nombres[0] in quedan and nombres[1] in quedan          # hoy: todas
    assert nombres[3] in quedan and nombres[2] not in quedan      # hace 10 días: sólo la última
    assert nombres[4] not in quedan and nombres[5] not in quedan  # más de 30 días: se van


def test_texto_registro(con_datos):
    con = con_datos.db.lectura()
    d = datos_del_dia(con, dia_de(ahora_ms()))
    con.close()
    t = texto_registro(d)
    assert t.startswith("Registro de cambios") and "Marcos" in t


def test_restaurar_no_borra_lo_anotado_en_el_registro(con_datos, rutas):
    copias = Copias(con_datos.db, con_datos, rutas)
    copia = copias.crear("manual")
    con_datos.aplicar({"ops": [op("negocio", "_", "Cambio posterior")]}, "Marcos")
    con_datos.aplicar({"ops": [borrar("productos", "p1")]}, "Marcos")
    copias.restaurar(copia.name, "Marcos")
    textos = [r[0] for r in con_datos.db.con.execute("SELECT resumen FROM auditoria ORDER BY id")]
    assert any("Producto eliminado: Coca-Cola" in t for t in textos), "el borrado hecho después de la copia sigue en el registro"
    assert any("Se restauró la copia" in t for t in textos)
    ids = [r[0] for r in con_datos.db.con.execute("SELECT id FROM auditoria ORDER BY id")]
    assert ids == sorted(set(ids))


def test_texto_raro_en_el_excel_no_lo_rompe(con_datos, rutas):
    from predio.util import ahora_ms, dia_de

    ts = ahora_ms()
    v = venta("v9", 9, ts - 1000, 500)
    v["cliente"] = "=1+1"
    v["nota"] = "con\x0bcaracter ilegal"
    con_datos.aplicar({"ops": [op("ventas", "v9", v)]}, "Marcos")
    ad = ArchivosDiarios(con_datos.db, rutas)
    ad.generar_dia(dia_de(ts))
    ws = load_workbook(ad.ruta_excel(dia_de(ts)))["Ventas"]
    celdas = [c.value for f in ws.iter_rows(min_row=2) for c in f]
    assert "=1+1" in celdas and any("caracter ilegal" in str(c) for c in celdas)


def test_un_cambio_durante_la_escritura_no_se_pierde(con_datos, rutas):
    from predio.util import ahora_ms, dia_de

    dia = dia_de(ahora_ms())
    ad = ArchivosDiarios(con_datos.db, rutas)
    ad.marcar([dia])
    original = ad.generar_dia

    def lento(d):
        ad.marcar([d])            # llega otra venta mientras se escribe
        return original(d)

    ad.generar_dia = lento
    ad.escribir_ahora()
    assert ad.pendientes() == [dia]
    ad.generar_dia = original
    ad.escribir_ahora()
    assert ad.pendientes() == []
